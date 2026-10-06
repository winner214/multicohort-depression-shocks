#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
MHAS primary analysis: incident shocks, Model 1/2/3, subgroup, interaction, cumulative-burden, and exploratory pathway analyses.

Run:
python MHAS.py --input H_MHAS_d.dta --outdir MHAS --min-baseline-age 50

"""
from __future__ import annotations

import argparse
import gc
import math
import subprocess
import tempfile
import time
import warnings
from contextlib import contextmanager
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import statsmodels.api as sm
import statsmodels.formula.api as smf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import chi2_contingency, f_oneway
from sklearn.experimental import enable_iterative_imputer  # noqa: F401
from sklearn.impute import IterativeImputer
from sklearn.linear_model import BayesianRidge
from statsmodels.genmod.cov_struct import Exchangeable
from tqdm.auto import tqdm

warnings.filterwarnings("ignore")

WAVE_YEAR_MAP = {1: 2001, 2: 2003, 3: 2012, 4: 2015, 5: 2018, 6: 2021}
ALL_WAVES = [1, 2, 3, 4, 5]
ANALYSIS_WAVES = [2, 3, 4, 5]
PATH_EXPOSURE_WAVES = [2, 3, 4]  # used to construct next-wave pathways; exposure rows are wave 2-4 after requiring next_wave
OUTCOME = "Depression_Risk"

IND_COMPONENTS = ["cancer", "stroke", "heart", "diabetes"]
SOC_COMPONENTS = ["marital_shock", "child_death", "spouse_cancer", "job_exit", "income_decline_shock", "living_alone_transition"]
PATH_MEDIATORS = ["self_rated_health", "adl", "social_activity_decline"]
SUMMARY_TERMS = ["OverallShock_main", "IndShock_main", "SocShock_main"]
COMPONENT_TERMS = IND_COMPONENTS + SOC_COMPONENTS

MODEL2_COVARIATES = ["age", "male", "education"]
MODEL3_COVARIATES = ["age", "male", "education", "smoking", "drinking"]
IMPUTE_COVARIATES = ["age", "male", "education", "smoking", "drinking"]
BINARY_COVARIATES = ["male", "smoking", "drinking"]
INCLUDE_WAVE_FIXED_EFFECT = True
MIN_BASELINE_AGE_DEFAULT = 50.0

CSV_FLOAT_FORMAT = "%.6g"
TQDM_KW = dict(mininterval=1.0, smoothing=0.0)

COHORT_NAME = "MHAS"
OUT_PREFIX = f"{COHORT_NAME}_"

OUT_VARIABLE_RULES = f"{OUT_PREFIX}00_variable_rules.csv"
OUT_SAMPLE_FLOW = f"{OUT_PREFIX}01_sample_flow.csv"
OUT_TABLE1 = f"{OUT_PREFIX}02_characteristics_by_wave.xlsx"
OUT_MAIN_SUMMARY = f"{OUT_PREFIX}03_main_summary_model123.csv"
OUT_MAIN_COMPONENT = f"{OUT_PREFIX}04_main_component_model123.csv"
OUT_SUBGROUP = f"{OUT_PREFIX}05_subgroup_summary_model3.csv"
OUT_PATH_SUMMARY = f"{OUT_PREFIX}06_path_results_summary.csv"
OUT_PATH_BOOT = f"{OUT_PREFIX}07_path_results_bootstrap_ci.csv"
OUT_META_READY = f"{OUT_PREFIX}08_meta_ready_model3.csv"
OUT_INTERACTION = f"{OUT_PREFIX}09_interaction_summary_model3.csv"
OUT_SHOCK_COUNT_DIST = f"{OUT_PREFIX}10_shock_count_distribution.csv"
OUT_SHOCK_COUNT_MODEL = f"{OUT_PREFIX}11_shock_count_model123.csv"
OUT_SHOCK_COUNT_META = f"{OUT_PREFIX}12_shock_count_meta_ready.csv"
OUT_NOTES = f"{OUT_PREFIX}99_analysis_notes.txt"
OUT_TIMING = f"{OUT_PREFIX}timing_stats.csv"


class StepTimer:
    def __init__(self):
        self.records = []

    @contextmanager
    def track(self, step: str):
        t0 = time.perf_counter()
        print(f"[Start] {step}", flush=True)
        try:
            yield
        finally:
            sec = time.perf_counter() - t0
            self.records.append({"step": step, "seconds": round(sec, 2), "minutes": round(sec / 60.0, 2)})
            print(f"[Done] {step} | {sec:.2f}s", flush=True)

    def to_frame(self):
        return pd.DataFrame(self.records)


def ensure_dir(path: Path):
    path.mkdir(parents=True, exist_ok=True)


def fast_to_csv(df: pd.DataFrame, path: Path):
    df.to_csv(path, index=False, float_format=CSV_FLOAT_FORMAT)


def write_text(path: Path, text: str):
    path.write_text(text, encoding="utf-8")


def safe_numeric(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s, errors="coerce")


def binary_from_numeric(s: pd.Series) -> pd.Series:
    s = safe_numeric(s)
    return pd.Series(np.where(s.notna(), np.where(s >= 0.5, 1.0, 0.0), np.nan), index=s.index, dtype="float32")


def living_alone_from_hhres(s: pd.Series) -> pd.Series:
    """Construct living-alone status from household-size variable.

    The harmonized household-size variable counts all co-resident household
    members and includes the respondent. Therefore, household size == 1 denotes
    living alone; household size >= 2 denotes not living alone. Missing or
    special missing values remain missing.
    """
    x = safe_numeric(s)
    out = pd.Series(np.nan, index=s.index, dtype="float32")
    out.loc[x == 1] = 1.0
    out.loc[x >= 2] = 0.0
    return out.astype("float32")


def row_any_shock(df: pd.DataFrame, cols: List[str]) -> pd.Series:
    """Composite shock indicator using an observed-component rule.

    Final manuscript rule:
    - If any observed component equals 1, the composite shock is 1.
    - If at least one component is observed and no observed component equals 1,
      the composite shock is 0.
    - Only rows with all components missing remain missing.

    This is equivalent to a defensible ``skipna=True`` construction with an
    all-missing safeguard, and avoids geometric sample loss from single-item
    nonresponse in multi-component shock domains.
    """
    x = df[cols].apply(pd.to_numeric, errors="coerce")
    out = pd.Series(np.nan, index=df.index, dtype="float32")
    any_observed = x.notna().any(axis=1)
    any_one = x.eq(1).any(axis=1)
    out.loc[any_observed] = 0.0
    out.loc[any_one] = 1.0
    return out.astype("float32")


def cast_small_dtypes(df: pd.DataFrame) -> pd.DataFrame:
    for c in df.columns:
        if pd.api.types.is_numeric_dtype(df[c]):
            try:
                df[c] = pd.to_numeric(df[c], downcast="float")
            except Exception:
                pass
    return df


def existing_columns(path: Path) -> List[str]:
    rdr = pd.read_stata(path, iterator=True)
    return list(rdr.variable_labels().keys())


def resolve_col(colmap_lower: Dict[str, str], candidates: List[str], required: bool = True) -> Optional[str]:
    for c in candidates:
        if c.lower() in colmap_lower:
            return colmap_lower[c.lower()]
    if required:
        raise KeyError(f"Cannot find any of candidate columns: {candidates}")
    return None


def build_raw_column_map(path: Path) -> Dict[str, object]:
    colmap_lower = {c.lower(): c for c in existing_columns(path)}
    m: Dict[str, object] = {
        "id": resolve_col(colmap_lower, ["unhhidnp", "UNHHIDNP", "rahhidnp", "RAHHIDNP"]),
        "gender": resolve_col(colmap_lower, ["ragender", "RAGENDER"]),
        "education": resolve_col(colmap_lower, ["raeducl", "RAEDUCL"]),
    }
    waves = {}
    for w in ALL_WAVES:
        waves[w] = {
            "inw": resolve_col(colmap_lower, [f"inw{w}", f"INW{w}"]),
            "age": resolve_col(colmap_lower, [f"r{w}agey", f"R{w}AGEY"]),
            "mstat": resolve_col(colmap_lower, [f"r{w}mstat", f"R{w}MSTAT"]),
            "work": resolve_col(colmap_lower, [f"r{w}work", f"R{w}WORK"]),
            "cancer_state": resolve_col(colmap_lower, [f"r{w}cancre", f"R{w}CANCRE"]),
            "spouse_cancer_state": resolve_col(colmap_lower, [f"s{w}cancre", f"S{w}CANCRE"], required=False),
            "child_death_count": resolve_col(colmap_lower, [f"h{w}dchild", f"H{w}DCHILD"], required=False),
            "stroke_state": resolve_col(colmap_lower, [f"r{w}stroke", f"R{w}STROKE"]),
            "heart_state": resolve_col(colmap_lower, [f"r{w}hearte", f"R{w}HEARTE", f"r{w}hrtatte", f"R{w}HRTATTE"], required=False),
            "diabetes_state": resolve_col(colmap_lower, [f"r{w}diabe", f"R{w}DIABE"]),
            "self_rated_health": resolve_col(colmap_lower, [f"r{w}shlt", f"R{w}SHLT", f"r{w}shlta", f"R{w}SHLTA"]),
            "smoking": resolve_col(colmap_lower, [f"r{w}smoken", f"R{w}SMOKEN", f"r{w}smokev", f"R{w}SMOKEV"]),
            "drinking": resolve_col(colmap_lower, [f"r{w}drink", f"R{w}DRINK", f"r{w}drinkl", f"R{w}DRINKL", f"r{w}drinkev", f"R{w}DRINKEV"]),
            "adl_raw": resolve_col(colmap_lower, [f"r{w}adltot6", f"R{w}ADLTOT6", f"r{w}adlab_c", f"R{w}ADLAB_C"]),
            "cesd_score": resolve_col(colmap_lower, [f"r{w}cesd_m", f"R{w}CESD_M", f"r{w}cesdm_m", f"R{w}CESDM_M"]),
            "social_activity": resolve_col(colmap_lower, [f"r{w}socmn", f"R{w}SOCMN", f"r{w}socwk", f"R{w}SOCWK"], required=False),
            "income": resolve_col(colmap_lower, [f"h{w}itot", f"H{w}ITOT", f"hh{w}itot", f"HH{w}ITOT"]),
            "hhres": resolve_col(colmap_lower, [f"h{w}hhres", f"H{w}HHRES", f"hh{w}hhres", f"HH{w}HHRES"], required=True),
        }
    m["waves"] = waves
    return m


def required_raw_columns(path: Path) -> Tuple[List[str], Dict[str, object]]:
    m = build_raw_column_map(path)
    cols = [m["id"], m["gender"], m["education"]]
    for w in ALL_WAVES:
        for v in m["waves"][w].values():
            if v is not None:
                cols.append(v)
    return sorted(set(cols)), m


def make_wave_frame(df: pd.DataFrame, w: int, rawmap: Dict[str, object]) -> pd.DataFrame:
    wm = rawmap["waves"][w]
    mstat = safe_numeric(df[wm["mstat"]]).astype("float32")
    partnered_now = mstat.isin([1, 3])
    non_partner_now = mstat.notna() & (~partnered_now)

    spouse_raw = safe_numeric(df[wm["spouse_cancer_state"]]).astype("float32") if wm.get("spouse_cancer_state") is not None else pd.Series(np.nan, index=df.index, dtype="float32")
    spouse_cancer_state = pd.Series(np.nan, index=df.index, dtype="float32")
    spouse_cancer_state.loc[non_partner_now] = 0.0
    spouse_cancer_state.loc[partnered_now & spouse_raw.notna()] = spouse_raw.loc[partnered_now & spouse_raw.notna()]

    gender = safe_numeric(df[rawmap["gender"]]).astype("float32")
    male = pd.Series(np.nan, index=df.index, dtype="float32")
    male.loc[gender == 1] = 1.0
    male.loc[gender == 2] = 0.0

    edu = safe_numeric(df[rawmap["education"]]).astype("float32")
    edu = edu.where(edu.isin([1, 2, 3]), np.nan)

    adl_raw = safe_numeric(df[wm["adl_raw"]]).astype("float32")
    adl = pd.Series(np.nan, index=df.index, dtype="float32")
    adl.loc[adl_raw.notna()] = np.where(adl_raw.loc[adl_raw.notna()] > 0, 1.0, 0.0)

    soc_col = wm["social_activity"]
    social_activity = binary_from_numeric(df[soc_col]) if soc_col is not None else pd.Series(np.nan, index=df.index, dtype="float32")
    hhres_col = wm.get("hhres")
    household_size = safe_numeric(df[hhres_col]).astype("float32") if hhres_col is not None else pd.Series(np.nan, index=df.index, dtype="float32")
    living_alone = living_alone_from_hhres(df[hhres_col]) if hhres_col is not None else pd.Series(np.nan, index=df.index, dtype="float32")

    out = pd.DataFrame({
        "id": safe_numeric(df[rawmap["id"]]).astype("int64"),
        "wave": np.int16(w),
        "year": np.int16(WAVE_YEAR_MAP[w]),
        "inw": safe_numeric(df[wm["inw"]]).astype("float32"),
        "male": male,
        "age": safe_numeric(df[wm["age"]]).astype("float32"),
        "education": edu,
        "mstat_raw": mstat,
        "work": binary_from_numeric(df[wm["work"]]),
        "cancer_state": binary_from_numeric(df[wm["cancer_state"]]),
        "stroke_state": binary_from_numeric(df[wm["stroke_state"]]),
        "heart_state": binary_from_numeric(df[wm["heart_state"]]) if wm.get("heart_state") is not None else pd.Series(np.nan, index=df.index, dtype="float32"),
        "diabetes_state": binary_from_numeric(df[wm["diabetes_state"]]),
        "spouse_cancer_state": spouse_cancer_state.astype("float32"),
        "child_death_count": safe_numeric(df[wm["child_death_count"]]).astype("float32") if wm.get("child_death_count") is not None else pd.Series(np.nan, index=df.index, dtype="float32"),
        "self_rated_health": safe_numeric(df[wm["self_rated_health"]]).astype("float32"),
        "smoking": binary_from_numeric(df[wm["smoking"]]),
        "drinking": binary_from_numeric(df[wm["drinking"]]),
        "adl_score_raw": adl_raw,
        "adl": adl.astype("float32"),
        "social_activity": social_activity.astype("float32"),
        "household_income": safe_numeric(df[wm["income"]]).astype("float32"),
        "household_size": household_size.astype("float32"),
        "living_alone": living_alone.astype("float32"),
        "cesd_raw": safe_numeric(df[wm["cesd_score"]]).astype("float32"),
    })
    out["age_group"] = np.where(out["age"].notna(), np.where(out["age"] > 60, 1.0, 0.0), np.nan).astype("float32")
    out["marital_state"] = np.nan
    mask_m = out["mstat_raw"].notna()
    out.loc[mask_m, "marital_state"] = np.where(out.loc[mask_m, "mstat_raw"].isin([5, 7]), 1.0, 0.0)
    out["marital_state"] = out["marital_state"].astype("float32")
    out[OUTCOME] = np.nan
    mask_y = out["cesd_raw"].notna()
    out.loc[mask_y, OUTCOME] = np.where(out.loc[mask_y, "cesd_raw"] >= 5, 1.0, 0.0)
    out[OUTCOME] = out[OUTCOME].astype("float32")
    return out


def make_income_tertile_by_wave(long_df: pd.DataFrame) -> pd.Series:
    out = pd.Series(np.nan, index=long_df.index, dtype="float32")
    for w in sorted(long_df["wave"].dropna().unique()):
        idx = long_df.index[long_df["wave"] == w]
        s = pd.to_numeric(long_df.loc[idx, "household_income"], errors="coerce")
        valid = s.notna()
        if valid.sum() < 30 or s[valid].nunique() < 3:
            continue
        ranks = s[valid].rank(method="first")
        cats = pd.qcut(ranks, q=3, labels=[1, 2, 3])
        out.loc[idx[valid]] = cats.astype("float32").values
    return out.astype("float32")


def incident_transition(current: pd.Series, previous: pd.Series) -> pd.Series:
    out = pd.Series(np.nan, index=current.index, dtype="float32")
    mask = current.notna() & previous.notna()
    out.loc[mask] = np.where((previous.loc[mask] == 0.0) & (current.loc[mask] == 1.0), 1.0, 0.0)
    return out.astype("float32")


def decline_transition(current: pd.Series, previous: pd.Series) -> pd.Series:
    out = pd.Series(np.nan, index=current.index, dtype="float32")
    mask = current.notna() & previous.notna()
    out.loc[mask] = np.where(previous.loc[mask] > current.loc[mask], 1.0, 0.0)
    return out.astype("float32")


def count_increase_transition(current: pd.Series, previous: pd.Series) -> pd.Series:
    out = pd.Series(np.nan, index=current.index, dtype="float32")
    mask = current.notna() & previous.notna()
    out.loc[mask] = np.where(current.loc[mask] > previous.loc[mask], 1.0, 0.0)
    return out.astype("float32")


def build_long_panel(df_wide: pd.DataFrame, rawmap: Dict[str, object]) -> pd.DataFrame:
    frames = []
    for w in tqdm(ALL_WAVES, desc="Build waves", unit="wave", **TQDM_KW):
        frames.append(make_wave_frame(df_wide, w, rawmap))
    long_df = pd.concat(frames, axis=0, ignore_index=True)
    del frames
    gc.collect()

    long_df = long_df.sort_values(["id", "wave"], kind="mergesort").reset_index(drop=True)
    baseline_wave = min(ALL_WAVES)
    baseline_age_map = long_df.loc[long_df["wave"] == baseline_wave, ["id", "age"]].drop_duplicates("id").set_index("id")["age"]
    long_df["baseline_age"] = long_df["id"].map(baseline_age_map)
    if long_df["baseline_age"].isna().any():
        first_nonmissing_age = long_df.groupby("id", sort=False)["age"].transform("first")
        long_df["baseline_age"] = long_df["baseline_age"].fillna(first_nonmissing_age)
    long_df["baseline_age"] = pd.to_numeric(long_df["baseline_age"], errors="coerce").astype("float32")
    long_df["baseline_age_50plus"] = np.where(long_df["baseline_age"].notna(), np.where(long_df["baseline_age"] >= 50.0, 1.0, 0.0), np.nan).astype("float32")
    long_df["income_level_3"] = make_income_tertile_by_wave(long_df)


    state_cols = ["cancer_state", "stroke_state", "heart_state", "diabetes_state", "marital_state", "spouse_cancer_state", "child_death_count", "work", "income_level_3", "social_activity", "living_alone"]
    for c in state_cols:
        long_df[f"prev_{c}"] = long_df.groupby("id", sort=False)[c].shift(1).astype("float32")

    long_df["cancer"] = incident_transition(long_df["cancer_state"], long_df["prev_cancer_state"])
    long_df["stroke"] = incident_transition(long_df["stroke_state"], long_df["prev_stroke_state"])
    long_df["heart"] = incident_transition(long_df["heart_state"], long_df["prev_heart_state"])
    long_df["diabetes"] = incident_transition(long_df["diabetes_state"], long_df["prev_diabetes_state"])
    long_df["marital_shock"] = incident_transition(long_df["marital_state"], long_df["prev_marital_state"])
    long_df["spouse_cancer"] = incident_transition(long_df["spouse_cancer_state"], long_df["prev_spouse_cancer_state"])
    long_df["child_death"] = count_increase_transition(long_df["child_death_count"], long_df["prev_child_death_count"])

    long_df["job_exit"] = np.nan
    mask_job = long_df["prev_work"].notna() & long_df["work"].notna()
    long_df.loc[mask_job, "job_exit"] = np.where((long_df.loc[mask_job, "prev_work"] == 1.0) & (long_df.loc[mask_job, "work"] == 0.0), 1.0, 0.0)
    long_df["job_exit"] = long_df["job_exit"].astype("float32")

    long_df["income_decline_shock"] = decline_transition(long_df["income_level_3"], long_df["prev_income_level_3"])

    long_df["social_activity_decline"] = np.nan
    mask_soc = long_df["prev_social_activity"].notna() & long_df["social_activity"].notna()
    long_df.loc[mask_soc, "social_activity_decline"] = np.where((long_df.loc[mask_soc, "prev_social_activity"] == 1.0) & (long_df.loc[mask_soc, "social_activity"] == 0.0), 1.0, 0.0)
    long_df["social_activity_decline"] = long_df["social_activity_decline"].astype("float32")

    long_df["living_alone_transition"] = np.nan
    mask_live = long_df["prev_living_alone"].notna() & long_df["living_alone"].notna()
    long_df.loc[mask_live, "living_alone_transition"] = np.where((long_df.loc[mask_live, "prev_living_alone"] == 0.0) & (long_df.loc[mask_live, "living_alone"] == 1.0), 1.0, 0.0)
    long_df["living_alone_transition"] = long_df["living_alone_transition"].astype("float32")

    for c in COMPONENT_TERMS + ["social_activity_decline", "income_decline_shock"]:
        long_df.loc[long_df["wave"] == 1, c] = np.nan

    long_df["IndShock_main"] = row_any_shock(long_df, IND_COMPONENTS)
    long_df["SocShock_main"] = row_any_shock(long_df, SOC_COMPONENTS)
    long_df["OverallShock_main"] = row_any_shock(long_df, ["IndShock_main", "SocShock_main"])
    return long_df


def variable_rules() -> pd.DataFrame:
    rows = [
        ("Depression_Risk", "rWcesd_m", "1 if MHAS CES-D score >= 5; 0 otherwise"),
        ("baseline_age", "r1agey or first nonmissing rWagey", "baseline/entry age; participants with baseline_age < 50 are excluded from main, subgroup, interaction, and pathway analyses"),
        ("age", "rWagey", "current wave age after baseline-age restriction"),
        ("age_group", "rWagey", "0 if current age <= 60; 1 if current age > 60"),
        ("male", "ragender", "1 male; 0 female"),
        ("education", "raeducl", "time-invariant harmonized education copied to every wave; 1 low, 2 middle, 3 high"),
        ("smoking", "rWsmoken preferred; rWsmokev fallback", "binary current smoking indicator where available; fallback to ever smoking only if current smoking is unavailable"),
        ("drinking", "rWdrink", "binary drinking indicator"),
        ("self_rated_health", "rWshlta/rWshlt", "pathway variable; not a main covariate"),
        ("adl", "rWadlab_c or rWadltot6", "pathway variable; 1 if ADL difficulty count > 0"),
        ("social_activity", "rWsocmn preferred; rWsocwk fallback", "1 if monthly social activity; only available from wave3 onward in MHAS"),
        ("social_activity_decline", "rWsocmn/rWsocwk", "1 if previous social_activity=1 and current social_activity=0"),
        ("household_income", "hWitot", "continuous household total income"),
        ("household_size", "hWhhres/HwHHRES", "number of household residents including the respondent"),
        ("living_alone", "hWhhres/HwHHRES", "1 if household_size == 1; 0 if household_size >= 2"),
        ("living_alone_transition", "hWhhres/HwHHRES", "incident social-isolation shock: previous living_alone=0 and current living_alone=1"),
        ("income_level_3", "hWitot", "wave-specific tertile: 1 low, 2 middle, 3 high"),
        ("income_decline_shock", "income_level_3", "incident socioeconomic shock: 1 if previous income tertile > current income tertile"),
        ("cancer", "rWcancre", "incident shock: previous=0 and current=1"),
        ("stroke", "rWstroke", "incident shock: previous=0 and current=1"),
        ("heart", "rWhearte preferred; rWhrtatte fallback", "incident shock: previous=0 and current=1"),
        ("diabetes", "rWdiabe", "incident shock: previous=0 and current=1"),
        ("marital_shock", "rWmstat", "previous not divorced/widowed and current divorced/widowed"),
        ("child_death", "hWdchild", "incident shock: current deceased-child count > previous deceased-child count"),
        ("spouse_cancer", "sWcancre", "previous spouse cancer=0 and current=1; no partner coded as 0"),
        ("job_exit", "rWwork", "previous paid work=1 and current paid work=0"),
        ("IndShock_main", "incident components", "row-wise any observed shock among cancer, stroke, heart, and diabetes"),
        ("SocShock_main", "incident components", "row-wise any observed shock among marital_shock, child_death, spouse_cancer, job_exit, income_decline_shock, and living_alone_transition"),
        ("OverallShock_main", "incident components", "max(IndShock_main, SocShock_main)"),
        ("shock_count_all", "incident components", "sum of observed incident shock components across individual and social-disruption shocks"),
        ("shock_count_cat", "shock_count_all", "0, 1, 2, and 3 for three or more shocks; used for dose-response models"),
        ("shock_count_plot", "shock_count_all", "0, 1, 2, 3, and 4 for four or more shocks; used for frequency distribution plots"),
    ]
    return pd.DataFrame(rows, columns=["analysis_variable", "raw_variable", "definition"])


def sample_flow_table(long_df: pd.DataFrame, analysis_df: pd.DataFrame, path_df: pd.DataFrame, min_baseline_age: float = MIN_BASELINE_AGE_DEFAULT) -> pd.DataFrame:
    rows = []
    rows.append({"step": "wide_persons_input", "n": int(long_df["id"].nunique())})
    rows.append({"step": "long_rows_wave1_to_wave5", "n": int(len(long_df))})
    rows.append({"step": f"persons_baseline_age_ge_{int(min_baseline_age)}", "n": int(long_df.loc[long_df["baseline_age"] >= min_baseline_age, "id"].nunique())})
    tmp = long_df[long_df["wave"].isin(ANALYSIS_WAVES)]
    rows.append({"step": "rows_wave2_to_wave5_for_incident_shock_analysis", "n": int(len(tmp))})
    tmp = tmp[tmp["baseline_age"] >= min_baseline_age]
    rows.append({"step": f"rows_wave2_to_wave5_baseline_age_ge_{int(min_baseline_age)}", "n": int(len(tmp))})
    tmp = tmp[tmp["inw"] == 1]
    rows.append({"step": "rows_current_wave_responded", "n": int(len(tmp))})
    tmp = tmp[tmp[OUTCOME].notna()]
    rows.append({"step": "rows_nonmissing_outcome", "n": int(len(tmp))})
    tmp2 = tmp.dropna(subset=SUMMARY_TERMS + COMPONENT_TERMS)
    rows.append({"step": "rows_complete_all_main_shocks_for_reference", "n": int(len(tmp2))})
    rows.append({"step": "rows_in_main_model_base_dataset_outcome_nonmissing", "n": int(len(analysis_df))})
    rows.append({"step": "persons_in_main_model_dataset", "n": int(analysis_df["id"].nunique())})
    rows.append({"step": "rows_in_path_dataset", "n": int(len(path_df))})
    rows.append({"step": "persons_in_path_dataset", "n": int(path_df["id"].nunique()) if len(path_df) else 0})
    return pd.DataFrame(rows)


def _binary_pvalue(df: pd.DataFrame, var: str) -> float:
    tmp = df[["wave", var]].copy()
    tmp[var] = pd.to_numeric(tmp[var], errors="coerce")
    tmp = tmp[tmp[var].notna()].copy()
    if tmp.empty:
        return np.nan
    tab = pd.crosstab(tmp["wave"], tmp[var])
    for col in [0.0, 1.0]:
        if col not in tab.columns:
            tab[col] = 0
    if tab.shape[0] < 2:
        return np.nan
    try:
        _, p, _, _ = chi2_contingency(tab[[0.0, 1.0]].values)
        return float(p)
    except Exception:
        return np.nan


def _cont_pvalue(df: pd.DataFrame, var: str) -> float:
    groups = []
    for w in sorted(df["wave"].dropna().unique()):
        s = pd.to_numeric(df.loc[df["wave"] == w, var], errors="coerce").dropna()
        if len(s) > 0:
            groups.append(s.values)
    if len(groups) < 2:
        return np.nan
    try:
        _, p = f_oneway(*groups)
        return float(p)
    except Exception:
        return np.nan


def export_characteristics_xlsx(df: pd.DataFrame, out_path: Path) -> None:
    specs = [
        ("Outcome", None, None),
        ("Depression_Risk, n (%)", "Depression_Risk", "binary"),

        ("Individual shock components", None, None),
        ("cancer, n (%)", "cancer", "binary"),
        ("stroke, n (%)", "stroke", "binary"),
        ("heart disease, n (%)", "heart", "binary"),
        ("diabetes, n (%)", "diabetes", "binary"),

        ("Social shock components", None, None),
        ("marital shock, n (%)", "marital_shock", "binary"),
        ("child death, n (%)", "child_death", "binary"),
        ("spouse cancer, n (%)", "spouse_cancer", "binary"),
        ("job exit, n (%)", "job_exit", "binary"),
        ("income decline, n (%)", "income_decline_shock", "binary"),
        ("transition to living alone, n (%)", "living_alone_transition", "binary"),

        ("Covariates", None, None),
        ("baseline_age, mean (SD)", "baseline_age", "cont"),
        ("age, mean (SD)", "age", "cont"),
        ("age > 60, n (%)", "age_group", "binary"),
        ("male, n (%)", "male", "binary"),
        ("education, mean (SD)", "education", "cont"),
        ("smoking, n (%)", "smoking", "binary"),
        ("drinking, n (%)", "drinking", "binary"),

        ("Pathway variables", None, None),
        ("self_rated_health, mean (SD)", "self_rated_health", "cont"),
        ("ADL limitation, n (%)", "adl", "binary"),
        ("social_activity, n (%)", "social_activity", "binary"),
        ("social_activity_decline, n (%)", "social_activity_decline", "binary"),
    ]
    table_waves = ANALYSIS_WAVES
    wave_ns = {int(w): int((df["wave"] == w).sum()) for w in table_waves}
    overall_n = int(len(df))
    wb = Workbook()
    ws = wb.active
    ws.title = "Characteristics"
    ws.append(["Characteristics"] + [f"Wave{w}\n(N={wave_ns.get(w,0)})" for w in table_waves] + [f"Overall\n(N={overall_n})", "P-Value"])
    section_rows = []
    for label, var, vtype in specs:
        if var is None:
            ws.append([label] + [""] * (len(table_waves) + 2))
            section_rows.append(ws.max_row)
            continue
        row = [label]
        for w in table_waves:
            s = pd.to_numeric(df.loc[df["wave"] == w, var], errors="coerce")
            if vtype == "binary":
                row.append("NA" if s.notna().sum() == 0 else f"{int((s == 1).sum())} ({(int((s == 1).sum())/s.notna().sum()*100):.1f})")
            else:
                row.append("NA" if s.notna().sum() == 0 else f"{s.mean():.1f} ({s.std(ddof=1):.1f})")
        s_all = pd.to_numeric(df[var], errors="coerce")
        if vtype == "binary":
            row.append("NA" if s_all.notna().sum() == 0 else f"{int((s_all == 1).sum())} ({(int((s_all == 1).sum())/s_all.notna().sum()*100):.1f})")
            p = _binary_pvalue(df, var)
        else:
            row.append("NA" if s_all.notna().sum() == 0 else f"{s_all.mean():.1f} ({s_all.std(ddof=1):.1f})")
            p = _cont_pvalue(df, var)
        row.append("" if pd.isna(p) else ("<0.001" if p < 0.001 else f"{p:.3f}"))
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    thin_gray = Side(style="thin", color="BFBFBF")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
            if r in section_rows:
                cell.fill = section_fill
                cell.font = Font(bold=True)
    ws.column_dimensions[get_column_letter(1)].width = 36
    for c in range(2, ws.max_column):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.column_dimensions[get_column_letter(ws.max_column)].width = 10
    ws.freeze_panes = "B2"
    wb.save(out_path)


def prepare_main_dataset(long_df: pd.DataFrame, min_baseline_age: float = MIN_BASELINE_AGE_DEFAULT) -> pd.DataFrame:
    """Main-analysis base dataset.

    Missing handling follows the specified analysis plan:
    - outcome missing: excluded here;
    - shock/exposure missing: NOT globally excluded here, but deleted inside
      each model for the specific exposure being fitted;
    - income decline and transition to living alone are included in main social-disruption shock analysis.
    """
    ana = long_df[long_df["wave"].isin(ANALYSIS_WAVES)].copy()
    ana = ana[ana["baseline_age"].notna() & (ana["baseline_age"] >= min_baseline_age)].copy()
    ana = ana[ana["inw"] == 1].copy()
    ana = ana[ana[OUTCOME].notna()].copy()
    return ana.sort_values(["id", "wave"], kind="mergesort").reset_index(drop=True)


def impute_covariates(base_df: pd.DataFrame, m: int, seed: int, mi_max_iter: int, mi_n_nearest_features: Optional[int], timer: StepTimer) -> List[pd.DataFrame]:
    out = []
    cov_df = base_df[IMPUTE_COVARIATES].copy().astype("float32")
    posterior = m > 1
    with timer.track("multiple_imputation_covariates"):
        for i in tqdm(range(m), desc="MI-covariates", unit="imp", **TQDM_KW):
            imp = IterativeImputer(estimator=BayesianRidge(), max_iter=mi_max_iter, n_nearest_features=mi_n_nearest_features, sample_posterior=posterior, random_state=seed + i, skip_complete=True, initial_strategy="most_frequent")
            arr = imp.fit_transform(cov_df)
            imp_df = base_df.copy()
            imp_cov = pd.DataFrame(arr, columns=IMPUTE_COVARIATES, index=base_df.index)
            for col in BINARY_COVARIATES:
                imp_cov[col] = np.where(imp_cov[col] >= 0.5, 1.0, 0.0)
            imp_cov["education"] = np.round(imp_cov["education"]).clip(1, 3)
            imp_cov["age"] = imp_cov["age"].clip(lower=0)
            for col in IMPUTE_COVARIATES:
                imp_df[col] = imp_cov[col].astype("float32")
            imp_df["age_group"] = np.where(imp_df["age"] > 60, 1.0, 0.0).astype("float32")
            out.append(imp_df)
            del arr, imp_df, imp_cov, imp
            gc.collect()
    return out


def covariate_terms_for_model(model_id: str, omit: Optional[str] = None) -> List[str]:
    if model_id == "Model 1":
        covs = []
    elif model_id == "Model 2":
        covs = ["age", "male", "factor(education)"]
    elif model_id == "Model 3":
        covs = ["age", "male", "factor(education)", "smoking", "drinking"]
    else:
        raise ValueError(model_id)
    if omit is None:
        return covs
    out = []
    for c in covs:
        if omit == "education" and c == "factor(education)":
            continue
        if c == omit:
            continue
        out.append(c)
    return out


def make_formula(exposure: str, model_id: str, omit: Optional[str] = None) -> str:
    terms = [exposure] + covariate_terms_for_model(model_id, omit=omit)
    if INCLUDE_WAVE_FIXED_EFFECT:
        terms.append("factor(wave)")
    return f"{OUTCOME} ~ {' + '.join(terms)} + (1|id)"


def _run_lme4_once(df: pd.DataFrame, formula: str, optimizer: str, maxfun: int, singular_tol: float, workdir: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    csv_in = workdir / "input.csv"
    csv_out = workdir / "coef.csv"
    csv_diag = workdir / "diag.csv"
    r_script = workdir / "fit_glmer.R"
    df.to_csv(csv_in, index=False)
    script = f'''suppressPackageStartupMessages(library(lme4))
d <- read.csv("{csv_in.as_posix()}")
d$id <- factor(d$id)
d$wave <- factor(d$wave)
if("education" %in% colnames(d)) d$education <- factor(d$education)
if("age" %in% colnames(d)) d$age <- d$age / 10.0
warn_msgs <- character()
fit <- withCallingHandlers(
  glmer(
    formula = {formula!r},
    data = d,
    family = binomial(link = "logit"),
    control = glmerControl(optimizer = {optimizer!r}, optCtrl = list(maxfun = {maxfun}), calc.derivs = FALSE)
  ),
  warning = function(w) {{ warn_msgs <<- c(warn_msgs, conditionMessage(w)); invokeRestart("muffleWarning") }}
)
co <- as.data.frame(coef(summary(fit)))
co$term <- rownames(co)
rownames(co) <- NULL
write.csv(co, "{csv_out.as_posix()}", row.names = FALSE)
conv_msgs <- character()
if (!is.null(fit@optinfo$conv$lme4$messages)) conv_msgs <- c(conv_msgs, fit@optinfo$conv$lme4$messages)
if (!is.null(fit@optinfo$conv$messages)) conv_msgs <- c(conv_msgs, fit@optinfo$conv$messages)
opt_code <- NA
if (!is.null(fit@optinfo$conv$opt)) opt_code <- suppressWarnings(as.numeric(fit@optinfo$conv$opt))
diag <- data.frame(optimizer = {optimizer!r}, maxfun = {maxfun}, fit_ok = 1,
  singular = as.integer(isSingular(fit, tol = {singular_tol})), conv_code = opt_code,
  n_warnings = length(unique(warn_msgs)), warnings = paste(unique(warn_msgs), collapse = " || "),
  n_conv_messages = length(unique(conv_msgs)), conv_messages = paste(unique(conv_msgs), collapse = " || "),
  logLik = suppressWarnings(as.numeric(logLik(fit))), AIC = suppressWarnings(AIC(fit)), BIC = suppressWarnings(BIC(fit)),
  stringsAsFactors = FALSE)
write.csv(diag, "{csv_diag.as_posix()}", row.names = FALSE)
'''
    r_script.write_text(script, encoding="utf-8")
    proc = subprocess.run(["Rscript", str(r_script)], check=True, capture_output=True, text=True)
    out = pd.read_csv(csv_out)
    diag = pd.read_csv(csv_diag)
    diag["stdout"] = proc.stdout.strip()
    diag["stderr"] = proc.stderr.strip()
    out.rename(columns={"Estimate": "coef", "Std. Error": "se", "z value": "z", "Pr(>|z|)": "p_value"}, inplace=True)
    # Age is divided by 10 in R for numerical stability; the returned age
    # coefficient is intentionally retained as the per-10-year effect for
    # manuscript reporting. Do not divide it again in Python.
    out["OR"] = np.exp(np.clip(out["coef"].astype(float), -30, 30))
    out["OR_CI_low"] = np.exp(np.clip(out["coef"].astype(float) - 1.96 * out["se"].astype(float), -30, 30))
    out["OR_CI_high"] = np.exp(np.clip(out["coef"].astype(float) + 1.96 * out["se"].astype(float), -30, 30))
    return out[["term", "coef", "se", "z", "p_value", "OR", "OR_CI_low", "OR_CI_high"]], diag


def _has_conv_problem(diag_row: pd.Series) -> bool:
    text = " ".join([str(diag_row.get("warnings", "") or ""), str(diag_row.get("conv_messages", "") or ""), str(diag_row.get("stderr", "") or "")]).lower()
    return any(p in text for p in ["failed to converge", "unable to evaluate scaled gradient", "degenerate", "hessian", "boundary (singular) fit"])


def fit_lme4_once(df: pd.DataFrame, formula: str, optimizer_seq: List[str], maxfun: int = 200000, singular_tol: float = 1e-4) -> Tuple[pd.DataFrame, pd.DataFrame]:
    cols = list(set([OUTCOME, "id", "wave"] + SUMMARY_TERMS + COMPONENT_TERMS + IMPUTE_COVARIATES + ["age_group", "prev_income_level_3", "prev_living_alone", "prev_social_activity", "shock_count_cat", "shock_count_trend"]))
    cols = [c for c in cols if c in df.columns]
    d = df[cols].copy()
    d[OUTCOME] = d[OUTCOME].astype(int)
    attempts = []
    best_out = None
    best_idx = None
    with tempfile.TemporaryDirectory(prefix="mhas_glmer_") as td:
        td_path = Path(td)
        for opt in optimizer_seq:
            try:
                out, diag = _run_lme4_once(d, formula=formula, optimizer=opt, maxfun=maxfun, singular_tol=singular_tol, workdir=td_path)
                diag["fit_ok"] = 1
                diag["selected"] = 0
                diag["conv_problem"] = int(_has_conv_problem(diag.iloc[0]))
                attempts.append(diag)
                if best_out is None:
                    best_out = out
                    best_idx = len(attempts) - 1
                singular = int(pd.to_numeric(diag.iloc[0].get("singular", 0), errors="coerce") or 0)
                conv_problem = int(diag.iloc[0]["conv_problem"])
                if singular == 0 and conv_problem == 0:
                    best_out = out
                    best_idx = len(attempts) - 1
                    break
            except subprocess.CalledProcessError as exc:
                attempts.append(pd.DataFrame([{"optimizer": opt, "maxfun": maxfun, "fit_ok": 0, "singular": np.nan, "conv_code": np.nan, "n_warnings": np.nan, "warnings": "", "n_conv_messages": np.nan, "conv_messages": "", "logLik": np.nan, "AIC": np.nan, "BIC": np.nan, "stdout": (exc.stdout or "").strip(), "stderr": (exc.stderr or "").strip(), "selected": 0, "conv_problem": 1, "error": str(exc)}]))
        if best_out is None:
            raise RuntimeError(f"All optimizers failed for formula: {formula}")
    diags = pd.concat(attempts, axis=0, ignore_index=True)
    if best_idx is not None:
        diags.loc[best_idx, "selected"] = 1
    del d
    gc.collect()
    return best_out, diags


def pool_rubin(results: List[pd.DataFrame]) -> pd.DataFrame:
    terms = sorted(set().union(*[set(r["term"].tolist()) for r in results]))
    m = len(results)
    rows = []
    for term in terms:
        qs, us = [], []
        for r in results:
            rr = r.loc[r["term"] == term]
            if rr.empty:
                continue
            qs.append(float(rr.iloc[0]["coef"]))
            us.append(float(rr.iloc[0]["se"]) ** 2)
        if not qs:
            continue
        q = np.asarray(qs, dtype=float)
        u = np.asarray(us, dtype=float)
        q_bar = float(np.mean(q))
        u_bar = float(np.mean(u))
        b = float(np.var(q, ddof=1)) if len(q) > 1 else 0.0
        t = u_bar + (1.0 + 1.0 / max(m, 1)) * b
        se = math.sqrt(max(t, 1e-12))
        z = q_bar / se if se > 0 else np.nan
        p = math.erfc(abs(float(z)) / math.sqrt(2.0)) if pd.notna(z) else np.nan
        rows.append({"term": term, "m": m, "coef": q_bar, "se": se, "z": z, "p_value": p, "within_var": u_bar, "between_var": b, "total_var": t, "OR": math.exp(np.clip(q_bar, -30, 30)), "OR_CI_low": math.exp(np.clip(q_bar - 1.96 * se, -30, 30)), "OR_CI_high": math.exp(np.clip(q_bar + 1.96 * se, -30, 30))})
    return pd.DataFrame(rows)


def fit_one_exposure_all_imputations(imputed_list: List[pd.DataFrame], exposure: str, model_id: str, timer: StepTimer, optimizer_seq: List[str], maxfun: int, singular_tol: float, omit: Optional[str] = None, label_prefix: str = "") -> Tuple[pd.DataFrame, pd.DataFrame]:
    formula = make_formula(exposure, model_id, omit=omit)
    coef_tables, diag_tables = [], []
    step_name = f"{label_prefix}{model_id}_{exposure}".replace(" ", "_")
    with timer.track(f"fit_{step_name}"):
        for k, d in enumerate(tqdm(imputed_list, desc=step_name, unit="imp", **TQDM_KW), start=1):
            dd = d.dropna(subset=[OUTCOME, exposure]).copy()
            if dd[exposure].nunique(dropna=True) < 2 or dd[OUTCOME].nunique(dropna=True) < 2:
                continue
            coef_df, diag_df = fit_lme4_once(dd, formula, optimizer_seq=optimizer_seq, maxfun=maxfun, singular_tol=singular_tol)
            coef_df["imputation"] = k
            diag_df["imputation"] = k
            diag_df["exposure"] = exposure
            diag_df["model"] = model_id
            diag_df["formula"] = formula
            coef_tables.append(coef_df)
            diag_tables.append(diag_df)
            gc.collect()
    if not coef_tables:
        empty = pd.DataFrame([{"term": exposure, "m": 0, "coef": np.nan, "se": np.nan, "z": np.nan, "p_value": np.nan, "OR": np.nan, "OR_CI_low": np.nan, "OR_CI_high": np.nan, "exposure": exposure, "model": model_id, "formula": formula}])
        return empty, pd.DataFrame()
    pooled = pool_rubin(coef_tables)
    keep = pooled[pooled["term"] == exposure].copy()
    keep["exposure"] = exposure
    keep["model"] = model_id
    keep["formula"] = formula
    diags = pd.concat(diag_tables, axis=0, ignore_index=True) if diag_tables else pd.DataFrame()
    return keep, diags


def fit_terms_model123(imputed_list: List[pd.DataFrame], exposures: List[str], timer: StepTimer, optimizer_seq: List[str], maxfun: int, singular_tol: float, label_prefix: str = "") -> Tuple[pd.DataFrame, pd.DataFrame]:
    rows, diags = [], []
    for model_id in ["Model 1", "Model 2", "Model 3"]:
        for exposure in exposures:
            res, diag = fit_one_exposure_all_imputations(imputed_list, exposure, model_id, timer, optimizer_seq=optimizer_seq, maxfun=maxfun, singular_tol=singular_tol, label_prefix=label_prefix)
            rows.append(res)
            if len(diag):
                diags.append(diag)
    return pd.concat(rows, axis=0, ignore_index=True), pd.concat(diags, axis=0, ignore_index=True) if diags else pd.DataFrame()



def add_shock_count_variables(df: pd.DataFrame) -> pd.DataFrame:
    """Add cumulative incident-shock burden variables for reviewer-requested analyses.

    shock_count_all counts observed incident shock components across the full
    component set. Missing component values are treated as 0 for this cumulative
    burden analysis to preserve the main analytical sample; the number of
    missing component values is exported in the distribution table for
    transparency.
    """
    out = df.copy()
    available = [c for c in COMPONENT_TERMS if c in out.columns]
    if not available:
        out["shock_count_all"] = np.nan
        out["shock_count_cat"] = np.nan
        out["shock_count_trend"] = np.nan
        out["shock_count_plot"] = np.nan
        out["shock_count_missing_components"] = np.nan
        return out
    comp = out[available].apply(pd.to_numeric, errors="coerce")
    out["shock_count_missing_components"] = comp.isna().sum(axis=1).astype("float32")
    out["shock_count_all"] = comp.fillna(0).sum(axis=1).astype("float32")
    out["shock_count_cat"] = out["shock_count_all"].clip(upper=3).astype("float32")
    out["shock_count_trend"] = out["shock_count_cat"].astype("float32")
    out["shock_count_plot"] = out["shock_count_all"].clip(upper=4).astype("float32")
    return out


def shock_count_distribution(df: pd.DataFrame) -> pd.DataFrame:
    """Export person-wave distribution of cumulative shock burden by wave and overall."""
    label_map = {
        0.0: "0 shocks",
        1.0: "1 shock",
        2.0: "2 shocks",
        3.0: "3 shocks",
        4.0: "4 or more shocks",
    }
    rows = []
    work = df.copy()
    work["shock_count_plot"] = pd.to_numeric(work["shock_count_plot"], errors="coerce")
    for wave_value, sub in [("Overall", work)] + [(int(w), work[work["wave"] == w]) for w in sorted(work["wave"].dropna().unique())]:
        denom = int(sub["shock_count_plot"].notna().sum())
        for level in [0.0, 1.0, 2.0, 3.0, 4.0]:
            n = int((sub["shock_count_plot"] == level).sum()) if denom else 0
            rows.append({
                "cohort": COHORT_NAME,
                "wave": wave_value,
                "shock_count_plot": int(level),
                "shock_count_label": label_map[level],
                "n_person_waves": n,
                "denominator_person_waves": denom,
                "percent_person_waves": (n / denom * 100.0) if denom else np.nan,
                "n_persons": int(sub.loc[sub["shock_count_plot"] == level, "id"].nunique()) if denom else 0,
                "mean_missing_components": float(pd.to_numeric(sub.get("shock_count_missing_components", pd.Series(dtype=float)), errors="coerce").mean()) if len(sub) else np.nan,
            })
    return pd.DataFrame(rows)


def make_shock_count_formula(model_id: str, trend: bool = False) -> str:
    exposure_term = "shock_count_trend" if trend else "factor(shock_count_cat)"
    terms = [exposure_term] + covariate_terms_for_model(model_id)
    if INCLUDE_WAVE_FIXED_EFFECT:
        terms.append("factor(wave)")
    return f"{OUTCOME} ~ {' + '.join(terms)} + (1|id)"


def _shock_count_level_from_term(term: str) -> Tuple[object, str]:
    t = str(term)
    if t == "shock_count_trend":
        return "trend", "Per additional shock category"
    if "factor(shock_count_cat)" in t:
        level = t.replace("factor(shock_count_cat)", "")
        try:
            level_int = int(float(level))
        except Exception:
            level_int = level
        label = {1: "1 shock vs 0 shocks", 2: "2 shocks vs 0 shocks", 3: "3 or more shocks vs 0 shocks"}.get(level_int, f"{level} vs 0 shocks")
        return level_int, label
    return np.nan, t


def fit_shock_count_model123(
    imputed_list: List[pd.DataFrame],
    timer: StepTimer,
    optimizer_seq: List[str],
    maxfun: int,
    singular_tol: float,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Fit categorical and trend models for cumulative shock burden under Model 1/2/3."""
    rows, diags = [], []
    for model_id in ["Model 1", "Model 2", "Model 3"]:
        for trend in [False, True]:
            formula = make_shock_count_formula(model_id, trend=trend)
            coef_tables, diag_tables = [], []
            needed_exposure = "shock_count_trend" if trend else "shock_count_cat"
            step_name = f"shock_count_{'trend' if trend else 'category'}_{model_id}".replace(" ", "_")
            with timer.track(f"fit_{step_name}"):
                for k, d in enumerate(tqdm(imputed_list, desc=step_name, unit="imp", **TQDM_KW), start=1):
                    dd = d.dropna(subset=[OUTCOME, needed_exposure]).copy()
                    if dd[needed_exposure].nunique(dropna=True) < 2 or dd[OUTCOME].nunique(dropna=True) < 2:
                        continue
                    coef_df, diag_df = fit_lme4_once(dd, formula, optimizer_seq=optimizer_seq, maxfun=maxfun, singular_tol=singular_tol)
                    coef_df["imputation"] = k
                    diag_df["imputation"] = k
                    diag_df["exposure"] = "shock_count_trend" if trend else "shock_count_cat"
                    diag_df["model"] = model_id
                    diag_df["formula"] = formula
                    coef_tables.append(coef_df)
                    diag_tables.append(diag_df)
                    gc.collect()
            if not coef_tables:
                continue
            pooled = pool_rubin(coef_tables)
            if trend:
                keep = pooled[pooled["term"] == "shock_count_trend"].copy()
            else:
                keep = pooled[pooled["term"].astype(str).str.startswith("factor(shock_count_cat)")].copy()
            if keep.empty:
                continue
            keep["analysis_type"] = "shock_count_trend" if trend else "shock_count_category"
            keep["exposure"] = "shock_count_trend" if trend else "shock_count_cat"
            keep["model"] = model_id
            keep["formula"] = formula
            parsed = keep["term"].apply(_shock_count_level_from_term)
            keep["shock_count_level"] = [x[0] for x in parsed]
            keep["contrast"] = [x[1] for x in parsed]
            rows.append(keep)
            if diag_tables:
                diags.append(pd.concat(diag_tables, axis=0, ignore_index=True))
    out = pd.concat(rows, axis=0, ignore_index=True) if rows else pd.DataFrame()
    diag_out = pd.concat(diags, axis=0, ignore_index=True) if diags else pd.DataFrame()
    return out, diag_out


def make_shock_count_meta_ready(shock_count_results: pd.DataFrame) -> pd.DataFrame:
    """Return Model 3 log(OR)/SE table for cumulative shock burden meta-analysis."""
    if shock_count_results is None or shock_count_results.empty:
        return pd.DataFrame()
    tmp = shock_count_results[shock_count_results["model"] == "Model 3"].copy()
    tmp.insert(0, "cohort", COHORT_NAME)
    keep_cols = [
        "cohort", "analysis_type", "exposure", "contrast", "shock_count_level", "term", "model",
        "coef", "se", "z", "p_value", "OR", "OR_CI_low", "OR_CI_high", "m", "formula"
    ]
    for c in keep_cols:
        if c not in tmp.columns:
            tmp[c] = np.nan
    return tmp[keep_cols]


def fit_subgroup_model3(imputed_list: List[pd.DataFrame], exposures: List[str], timer: StepTimer, optimizer_seq: List[str], maxfun: int, singular_tol: float) -> pd.DataFrame:
    subgroup_specs = [("male", {0.0: "Female", 1.0: "Male"}, "male"), ("age_group", {0.0: "Age<=60", 1.0: "Age>60"}, None), ("education", {1.0: "Education low", 2.0: "Education middle", 3.0: "Education high"}, "education"), ("smoking", {0.0: "Non-smoking", 1.0: "Smoking"}, "smoking"), ("drinking", {0.0: "Non-drinking", 1.0: "Drinking"}, "drinking"), ("prev_income_level_3", {1.0: "Economic low", 2.0: "Economic middle", 3.0: "Economic high"}, None), ("prev_social_activity", {0.0: "No pre-shock social activity", 1.0: "Any pre-shock social activity"}, None), ("prev_living_alone", {0.0: "Not living alone", 1.0: "Living alone"}, None)]
    rows = []
    for sg_var, levels, omit in subgroup_specs:
        for level_val, label in levels.items():
            sub_imputed = []
            n_rows_each, n_ids_each = [], []
            for d in imputed_list:
                sub = d[d[sg_var] == level_val].copy()
                sub_imputed.append(sub)
                n_rows_each.append(len(sub))
                n_ids_each.append(sub["id"].nunique() if len(sub) else 0)
            for exposure in exposures:
                res, _ = fit_one_exposure_all_imputations(sub_imputed, exposure, "Model 3", timer, optimizer_seq=optimizer_seq, maxfun=maxfun, singular_tol=singular_tol, omit=omit, label_prefix=f"subgroup_{sg_var}_{label}_")
                res["subgroup_variable"] = sg_var
                res["subgroup_level"] = label
                res["subgroup_value"] = level_val
                res["n_rows_mean_across_imputations"] = float(np.mean(n_rows_each))
                res["n_ids_mean_across_imputations"] = float(np.mean(n_ids_each))
                rows.append(res)
    return pd.concat(rows, axis=0, ignore_index=True) if rows else pd.DataFrame()




def subgroup_term_for_formula(subgroup_var: str) -> str:
    """Return the R formula term used for the subgroup variable."""
    if subgroup_var == "education":
        return "factor(education)"
    if subgroup_var == "prev_income_level_3":
        return "factor(prev_income_level_3)"
    return subgroup_var


def omit_covariate_for_interaction(subgroup_var: str) -> Optional[str]:
    """Avoid duplicating the subgroup main effect in Model 3 interaction models.

    Continuous age is retained even when the modifier is age_group to reduce
    residual age confounding within broad age strata.
    """
    if subgroup_var == "education":
        return "education"
    if subgroup_var == "age_group":
        return None
    if subgroup_var in ["prev_income_level_3", "prev_social_activity", "prev_living_alone"]:
        return None
    return subgroup_var


def make_interaction_formula(exposure: str, subgroup_var: str) -> str:
    """Model 3 interaction formula: exposure * subgroup + Model 3 covariates + wave FE + random intercept."""
    sg_term = subgroup_term_for_formula(subgroup_var)
    omit = omit_covariate_for_interaction(subgroup_var)
    terms = [f"{exposure} * {sg_term}"] + covariate_terms_for_model("Model 3", omit=omit)
    if INCLUDE_WAVE_FIXED_EFFECT:
        terms.append("factor(wave)")
    return f"{OUTCOME} ~ {' + '.join(terms)} + (1|id)"


def is_interaction_term(term: str, exposure: str, subgroup_var: str) -> bool:
    """Identify interaction coefficient rows from lme4 output."""
    term = str(term)
    if ":" not in term or exposure not in term:
        return False
    if subgroup_var == "education":
        return "education" in term
    return subgroup_var in term


def fit_one_interaction_all_imputations(
    imputed_list: List[pd.DataFrame],
    exposure: str,
    subgroup_var: str,
    timer: StepTimer,
    optimizer_seq: List[str],
    maxfun: int,
    singular_tol: float,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Fit one Model 3 exposure-by-subgroup interaction across imputations."""
    formula = make_interaction_formula(exposure, subgroup_var)
    coef_tables, diag_tables = [], []
    n_rows_each, n_ids_each = [], []
    step_name = f"interaction_Model3_{exposure}_by_{subgroup_var}"
    with timer.track(f"fit_{step_name}"):
        for k, d in enumerate(tqdm(imputed_list, desc=step_name, unit="imp", **TQDM_KW), start=1):
            dd = d.dropna(subset=[OUTCOME, exposure, subgroup_var]).copy()
            n_rows_each.append(len(dd))
            n_ids_each.append(dd["id"].nunique() if len(dd) else 0)
            if len(dd) == 0:
                continue
            if dd[exposure].nunique(dropna=True) < 2 or dd[subgroup_var].nunique(dropna=True) < 2 or dd[OUTCOME].nunique(dropna=True) < 2:
                continue
            coef_df, diag_df = fit_lme4_once(dd, formula, optimizer_seq=optimizer_seq, maxfun=maxfun, singular_tol=singular_tol)
            coef_df["imputation"] = k
            diag_df["imputation"] = k
            diag_df["exposure"] = exposure
            diag_df["model"] = "Model 3 interaction"
            diag_df["subgroup_variable"] = subgroup_var
            diag_df["formula"] = formula
            coef_tables.append(coef_df)
            diag_tables.append(diag_df)
            gc.collect()
    if not coef_tables:
        empty = pd.DataFrame([{
            "term": np.nan, "m": 0, "coef": np.nan, "se": np.nan, "z": np.nan, "p_value": np.nan,
            "OR": np.nan, "OR_CI_low": np.nan, "OR_CI_high": np.nan,
            "exposure": exposure, "subgroup_variable": subgroup_var,
            "interaction_term": np.nan, "model": "Model 3 interaction",
            "n_rows_mean_across_imputations": float(np.mean(n_rows_each)) if n_rows_each else np.nan,
            "n_ids_mean_across_imputations": float(np.mean(n_ids_each)) if n_ids_each else np.nan,
            "formula": formula,
        }])
        return empty, pd.DataFrame()
    pooled = pool_rubin(coef_tables)
    keep = pooled[pooled["term"].apply(lambda x: is_interaction_term(x, exposure, subgroup_var))].copy()
    if keep.empty:
        keep = pd.DataFrame([{
            "term": np.nan, "m": len(coef_tables), "coef": np.nan, "se": np.nan, "z": np.nan, "p_value": np.nan,
            "within_var": np.nan, "between_var": np.nan, "total_var": np.nan,
            "OR": np.nan, "OR_CI_low": np.nan, "OR_CI_high": np.nan,
        }])
    keep["exposure"] = exposure
    keep["subgroup_variable"] = subgroup_var
    keep["interaction_term"] = keep["term"]
    keep["model"] = "Model 3 interaction"
    keep["n_rows_mean_across_imputations"] = float(np.mean(n_rows_each)) if n_rows_each else np.nan
    keep["n_ids_mean_across_imputations"] = float(np.mean(n_ids_each)) if n_ids_each else np.nan
    keep["formula"] = formula
    diags = pd.concat(diag_tables, axis=0, ignore_index=True) if diag_tables else pd.DataFrame()
    return keep, diags


def fit_interaction_summary_model3(
    imputed_list: List[pd.DataFrame],
    exposures: List[str],
    timer: StepTimer,
    optimizer_seq: List[str],
    maxfun: int,
    singular_tol: float,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Formal exposure-by-subgroup interaction tests for summary shocks under Model 3."""
    subgroup_vars = ["male", "age_group", "education", "smoking", "drinking", "prev_income_level_3", "prev_social_activity", "prev_living_alone"]
    rows, diags = [], []
    for subgroup_var in subgroup_vars:
        for exposure in exposures:
            res, diag = fit_one_interaction_all_imputations(
                imputed_list=imputed_list,
                exposure=exposure,
                subgroup_var=subgroup_var,
                timer=timer,
                optimizer_seq=optimizer_seq,
                maxfun=maxfun,
                singular_tol=singular_tol,
            )
            rows.append(res)
            if len(diag):
                diags.append(diag)
    out = pd.concat(rows, axis=0, ignore_index=True) if rows else pd.DataFrame()
    if not out.empty:
        # For binary subgroup variables, exp(coef) is the ratio of ORs comparing subgroup=1 vs subgroup=0.
        # For education, rows are contrast-specific ratio of ORs versus the low-education reference.
        out["interaction_OR_ratio"] = out["OR"]
        out["interaction_OR_ratio_CI_low"] = out["OR_CI_low"]
        out["interaction_OR_ratio_CI_high"] = out["OR_CI_high"]
    diag_out = pd.concat(diags, axis=0, ignore_index=True) if diags else pd.DataFrame()
    return out, diag_out

def prepare_path_dataset(long_df: pd.DataFrame, min_baseline_age: float = MIN_BASELINE_AGE_DEFAULT) -> pd.DataFrame:
    """Prepare strict adjacent-wave pathway dataset.

    Exposure is measured at wave t. Mediators and depressive risk are measured
    at the immediately following survey wave t+1. Next-wave variables are
    constructed before deleting rows with missing response/outcome values, so
    the code cannot accidentally connect wave t to the next non-missing but
    non-adjacent wave.
    """
    exposure_waves = sorted(list(PATH_EXPOSURE_WAVES))
    next_waves = sorted([w + 1 for w in exposure_waves if (w + 1) in set(ALL_WAVES)])
    path_panel_waves = sorted(set(exposure_waves + next_waves))

    path_df = long_df[long_df["wave"].isin(path_panel_waves)].copy()
    path_df = path_df[path_df["baseline_age"].notna() & (path_df["baseline_age"] >= min_baseline_age)].copy()
    path_df = path_df.sort_values(["id", "wave"], kind="mergesort").reset_index(drop=True)

    # Construct next-wave variables before filtering by response or outcome status.
    # This preserves the actual adjacent-wave structure and prevents wave skipping.
    for col in ["self_rated_health", "adl", "social_activity_decline", OUTCOME, "inw"]:
        path_df[f"next_{col}"] = path_df.groupby("id", sort=False)[col].shift(-1)

    path_df["next_wave"] = path_df.groupby("id", sort=False)["wave"].shift(-1)

    # Strictly allow only adjacent survey-wave transitions: t -> t+1.
    path_df = path_df[path_df["next_wave"] == path_df["wave"] + 1].copy()

    # Keep only the prespecified exposure waves after next-wave variables exist.
    path_df = path_df[path_df["wave"].isin(exposure_waves)].copy()

    # Require response and depressive-risk information at both current and next waves.
    # Current Depression_Risk is adjusted in the pathway models; next_Depression_Risk
    # is the pathway outcome.
    path_df = path_df[path_df["inw"] == 1].copy()
    path_df = path_df[path_df["next_inw"] == 1].copy()
    path_df = path_df[path_df[OUTCOME].notna()].copy()
    path_df = path_df[path_df[f"next_{OUTCOME}"].notna()].copy()

    path_df["next_wave"] = path_df["next_wave"].astype(int)
    return path_df.reset_index(drop=True)




def gee_fit(df: pd.DataFrame, formula: str, family: str, group_col: str = "id"):
    fam = sm.families.Binomial() if family == "binomial" else sm.families.Gaussian()
    model = smf.gee(formula=formula, groups=df[group_col], data=df, family=fam, cov_struct=Exchangeable())
    return model.fit()


def _predict_outcome_prob(out_res, base: pd.DataFrame, exposure: str, mediator_next: str, a_value: float, m_value) -> np.ndarray:
    """Predict Pr(next outcome=1) under assigned exposure and mediator values."""
    nd = base.copy()
    nd[exposure] = float(a_value)
    nd[mediator_next] = m_value
    pred = out_res.predict(nd)
    return np.asarray(pred, dtype=float)


def _counterfactual_mediation_rd(sub: pd.DataFrame, exposure: str, mediator_next: str, mediator_family: str, med_res, out_res) -> Dict[str, float]:
    """Parametric g-computation mediation effects on the risk-difference scale.

    Exposure is set to 0/1. For binary mediators, mediator distributions are
    integrated out. For continuous mediators, mediator conditional means are
    used. This avoids the invalid product-of-log-OR calculation for non-linear
    mediator/outcome models.
    """
    dm0 = sub.copy()
    dm1 = sub.copy()
    dm0[exposure] = 0.0
    dm1[exposure] = 1.0

    if mediator_family == "binomial":
        p_m0 = np.clip(np.asarray(med_res.predict(dm0), dtype=float), 1e-8, 1.0 - 1e-8)
        p_m1 = np.clip(np.asarray(med_res.predict(dm1), dtype=float), 1e-8, 1.0 - 1e-8)

        y_1_m1 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 1.0, 1.0)
        y_1_m0 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 1.0, 0.0)
        y_0_m1 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 0.0, 1.0)
        y_0_m0 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 0.0, 0.0)

        ey_1_mdist1 = p_m1 * y_1_m1 + (1.0 - p_m1) * y_1_m0
        ey_1_mdist0 = p_m0 * y_1_m1 + (1.0 - p_m0) * y_1_m0
        ey_0_mdist0 = p_m0 * y_0_m1 + (1.0 - p_m0) * y_0_m0
    else:
        m0 = np.asarray(med_res.predict(dm0), dtype=float)
        m1 = np.asarray(med_res.predict(dm1), dtype=float)

        ey_1_mdist1 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 1.0, m1)
        ey_1_mdist0 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 1.0, m0)
        ey_0_mdist0 = _predict_outcome_prob(out_res, sub, exposure, mediator_next, 0.0, m0)

    total_rd = float(np.nanmean(ey_1_mdist1 - ey_0_mdist0))
    indirect_rd = float(np.nanmean(ey_1_mdist1 - ey_1_mdist0))
    direct_rd = float(np.nanmean(ey_1_mdist0 - ey_0_mdist0))
    prop_rd = indirect_rd / total_rd if np.isfinite(total_rd) and abs(total_rd) > 1e-12 else np.nan
    prop_abs_rd = abs(indirect_rd) / (abs(indirect_rd) + abs(direct_rd)) if np.isfinite(indirect_rd) and np.isfinite(direct_rd) and (abs(indirect_rd) + abs(direct_rd)) > 1e-12 else np.nan

    return {
        "total_rd": total_rd,
        "direct_rd": direct_rd,
        "indirect_rd": indirect_rd,
        "prop_mediated_rd": prop_rd,
        "prop_mediated_abs_rd": prop_abs_rd,
    }


PATH_RESULT_METRICS = [
    "n_rows", "n_ids",
    "a_path", "b_path",
    "direct_logit", "indirect_logit", "total_logit",
    "direct_logit_ystd", "indirect_logit_ystd", "total_logit_ystd",
    "direct_or", "total_or",
    "total_rd", "direct_rd", "indirect_rd",
    "prop_mediated_raw", "prop_mediated_abs",
    "prop_mediated_rd", "prop_mediated_abs_rd",
]


def empty_path_result(path_name: str, exposure: str, mediator_label: str) -> Dict[str, float]:
    row: Dict[str, float] = {"path": path_name, "exposure": exposure, "mediator": mediator_label}
    for c in PATH_RESULT_METRICS:
        row[c] = np.nan
    row["method"] = "counterfactual_g_computation_rd_primary; coefficient_difference_secondary"
    row["product_method_used"] = 0
    return row


def fit_single_path(df: pd.DataFrame, path_name: str, exposure: str, mediator_next: str, mediator_family: str, mediator_label: str) -> Dict[str, float]:
    covs = ["age", "male", "education", "smoking", "drinking"]
    needed = ["id", "wave", "next_wave", OUTCOME, exposure, mediator_next] + covs + [f"next_{OUTCOME}"]
    sub = df[needed].copy().dropna()
    if sub.empty:
        raise ValueError(f"No usable rows for path {path_name}")

    med_formula = f"{mediator_next} ~ {exposure} + age + male + C(education) + smoking + drinking + {OUTCOME} + C(next_wave)"
    direct_formula = f"next_{OUTCOME} ~ {exposure} + {mediator_next} + age + male + C(education) + smoking + drinking + {OUTCOME} + C(next_wave)"
    total_formula = f"next_{OUTCOME} ~ {exposure} + age + male + C(education) + smoking + drinking + {OUTCOME} + C(next_wave)"

    med_res = gee_fit(sub, med_formula, mediator_family)
    direct_res = gee_fit(sub, direct_formula, "binomial")
    total_res = gee_fit(sub, total_formula, "binomial")

    a = float(med_res.params.get(exposure, np.nan))
    b = float(direct_res.params.get(mediator_next, np.nan))
    c_prime = float(direct_res.params.get(exposure, np.nan))
    c_total = float(total_res.params.get(exposure, np.nan))

    # Secondary coefficient-difference estimate. This replaces the invalid a*b
    # product on mixed/non-linear log-odds scales.
    indirect_logit = c_total - c_prime if np.isfinite(c_total) and np.isfinite(c_prime) else np.nan
    logit_y_scale = math.sqrt(math.pi ** 2 / 3.0)
    c_prime_ystd = c_prime / logit_y_scale if np.isfinite(c_prime) else np.nan
    c_total_ystd = c_total / logit_y_scale if np.isfinite(c_total) else np.nan
    indirect_ystd = c_total_ystd - c_prime_ystd if np.isfinite(c_total_ystd) and np.isfinite(c_prime_ystd) else np.nan

    prop_med_logit = indirect_logit / c_total if np.isfinite(c_total) and abs(c_total) > 1e-12 else np.nan
    prop_med_abs_logit = abs(indirect_logit) / (abs(indirect_logit) + abs(c_prime)) if np.isfinite(indirect_logit) and np.isfinite(c_prime) and (abs(indirect_logit) + abs(c_prime)) > 1e-12 else np.nan

    cf = _counterfactual_mediation_rd(sub, exposure, mediator_next, mediator_family, med_res, direct_res)

    return {
        "path": path_name,
        "exposure": exposure,
        "mediator": mediator_label,
        "n_rows": len(sub),
        "n_ids": sub["id"].nunique(),
        "a_path": a,
        "b_path": b,
        "direct_logit": c_prime,
        "indirect_logit": indirect_logit,
        "total_logit": c_total,
        "direct_logit_ystd": c_prime_ystd,
        "indirect_logit_ystd": indirect_ystd,
        "total_logit_ystd": c_total_ystd,
        "direct_or": math.exp(np.clip(c_prime, -30, 30)) if np.isfinite(c_prime) else np.nan,
        "total_or": math.exp(np.clip(c_total, -30, 30)) if np.isfinite(c_total) else np.nan,
        "total_rd": cf["total_rd"],
        "direct_rd": cf["direct_rd"],
        "indirect_rd": cf["indirect_rd"],
        "prop_mediated_raw": prop_med_logit,
        "prop_mediated_abs": prop_med_abs_logit,
        "prop_mediated_rd": cf["prop_mediated_rd"],
        "prop_mediated_abs_rd": cf["prop_mediated_abs_rd"],
        "method": "counterfactual_g_computation_rd_primary; coefficient_difference_secondary",
        "product_method_used": 0,
    }


def fit_path_models_mi(imputed_path_list: List[pd.DataFrame], path_specs: List[Tuple[str, str, str, str, str]], timer: StepTimer) -> pd.DataFrame:
    """Fit pathway models across covariate-imputed datasets and average estimates.

    Only covariates are imputed. Exposures, outcomes, and mediators remain
    complete-case within each path-specific model. Primary mediation estimates
    are counterfactual g-computation risk differences; coefficient-difference
    estimates are retained as secondary logit-scale diagnostics.
    """
    rows = []
    with timer.track("fit_path_models_point_estimates_mi"):
        for spec in path_specs:
            one_path = []
            for k, d in enumerate(imputed_path_list, start=1):
                try:
                    res = fit_single_path(d, *spec)
                    res["imputation"] = k
                    one_path.append(res)
                except Exception:
                    continue
            if not one_path:
                row = empty_path_result(spec[0], spec[1], spec[4])
                row["m"] = 0
                rows.append(row)
                continue
            tmp = pd.DataFrame(one_path)
            row = {"path": spec[0], "exposure": spec[1], "mediator": spec[4], "m": len(tmp)}
            for c in PATH_RESULT_METRICS:
                if c in tmp.columns:
                    row[c] = float(pd.to_numeric(tmp[c], errors="coerce").mean())
            row["method"] = "counterfactual_g_computation_rd_primary; coefficient_difference_secondary"
            row["product_method_used"] = 0
            rows.append(row)
    return pd.DataFrame(rows)


def bootstrap_paths_mi(imputed_path_list: List[pd.DataFrame], path_specs: List[Tuple[str, str, str, str, str]], n_boot: int, seed: int, timer: StepTimer) -> pd.DataFrame:
    """Bootstrap path models using rotating covariate-imputed datasets."""
    rng = np.random.default_rng(seed)
    if not imputed_path_list:
        return pd.DataFrame()
    ids = np.array(sorted(imputed_path_list[0]["id"].dropna().unique()))
    rows = []
    with timer.track("bootstrap_path_models_mi"):
        for b in tqdm(range(1, n_boot + 1), desc="Bootstrap-path-MI", unit="boot", **TQDM_KW):
            base_df = imputed_path_list[(b - 1) % len(imputed_path_list)]
            sample_ids = rng.choice(ids, size=len(ids), replace=True)
            boot_frames = []
            for i, sid in enumerate(sample_ids):
                tmp = base_df[base_df["id"] == sid].copy()
                tmp["id"] = i + 1
                boot_frames.append(tmp)
            boot_df = pd.concat(boot_frames, axis=0, ignore_index=True)
            for spec in path_specs:
                try:
                    res = fit_single_path(boot_df, *spec)
                    res["bootstrap"] = b
                    rows.append(res)
                except Exception:
                    row = empty_path_result(spec[0], spec[1], spec[4])
                    row["bootstrap"] = b
                    rows.append(row)
            del boot_frames, boot_df
            gc.collect()
    return pd.DataFrame(rows)


def bootstrap_paths(df: pd.DataFrame, path_specs: List[Tuple[str, str, str, str, str]], n_boot: int, seed: int, timer: StepTimer) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    ids = np.array(sorted(df["id"].dropna().unique()))
    rows = []
    with timer.track("bootstrap_path_models"):
        for b in tqdm(range(1, n_boot + 1), desc="Bootstrap-path", unit="boot", **TQDM_KW):
            sample_ids = rng.choice(ids, size=len(ids), replace=True)
            boot_frames = []
            for i, sid in enumerate(sample_ids):
                tmp = df[df["id"] == sid].copy()
                tmp["id"] = i + 1
                boot_frames.append(tmp)
            boot_df = pd.concat(boot_frames, axis=0, ignore_index=True)
            for spec in path_specs:
                try:
                    res = fit_single_path(boot_df, *spec)
                    res["bootstrap"] = b
                    rows.append(res)
                except Exception:
                    row = empty_path_result(spec[0], spec[1], spec[4])
                    row["bootstrap"] = b
                    rows.append(row)
            del boot_frames, boot_df
            gc.collect()
    return pd.DataFrame(rows)


def summarize_bootstrap(point_df: pd.DataFrame, boot_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    if boot_df is None or boot_df.empty:
        return point_df.copy()
    for path_name, sub in boot_df.groupby("path", dropna=False):
        pp = point_df[point_df["path"] == path_name]
        if pp.empty:
            continue
        point = pp.iloc[0].to_dict()
        row = point.copy()
        for metric in PATH_RESULT_METRICS:
            if metric not in sub.columns:
                row[f"{metric}_lcl"] = np.nan
                row[f"{metric}_ucl"] = np.nan
                continue
            s = pd.to_numeric(sub[metric], errors="coerce").dropna()
            if len(s) > 10:
                row[f"{metric}_lcl"] = float(np.quantile(s, 0.025))
                row[f"{metric}_ucl"] = float(np.quantile(s, 0.975))
            else:
                row[f"{metric}_lcl"] = np.nan
                row[f"{metric}_ucl"] = np.nan
        rows.append(row)
    return pd.DataFrame(rows)



def diag_problem(diag: pd.DataFrame) -> int:
    if diag is None or diag.empty:
        return 1
    return int(((diag["fit_ok"] != 1) | (diag["singular"].fillna(0) != 0) | (diag["conv_problem"] != 0)).any())


def make_meta_ready(summary_results: pd.DataFrame, component_results: pd.DataFrame) -> pd.DataFrame:
    """Return Model 3 log(OR)/SE table for cross-database random-effects meta-analysis."""
    rows = []
    for analysis_type, df in [("summary", summary_results), ("component", component_results)]:
        if df is None or df.empty:
            continue
        tmp = df[df["model"] == "Model 3"].copy()
        tmp["analysis_type"] = analysis_type
        keep_cols = [
            "analysis_type", "exposure", "term", "model", "coef", "se", "z", "p_value",
            "OR", "OR_CI_low", "OR_CI_high", "m", "formula"
        ]
        for c in keep_cols:
            if c not in tmp.columns:
                tmp[c] = np.nan
        rows.append(tmp[keep_cols])
    if not rows:
        return pd.DataFrame()
    out = pd.concat(rows, axis=0, ignore_index=True)
    out.insert(0, "cohort", COHORT_NAME)
    return out


def run_pipeline(args: argparse.Namespace) -> None:
    timer = StepTimer()
    input_path = Path(args.input).resolve()
    outdir = Path(args.outdir).resolve()
    ensure_dir(outdir)

    with timer.track("resolve_and_read_dta"):
        raw_cols, rawmap = required_raw_columns(input_path)
        raw = pd.read_stata(input_path, columns=raw_cols, convert_categoricals=False)
        raw = cast_small_dtypes(raw)

    with timer.track("build_long_panel"):
        long_df = build_long_panel(raw, rawmap=rawmap)
        del raw
        gc.collect()

    with timer.track("prepare_main_dataset"):
        ana = prepare_main_dataset(long_df, min_baseline_age=args.min_baseline_age)
        ana = add_shock_count_variables(ana)

    with timer.track("prepare_path_dataset"):
        path_df = prepare_path_dataset(long_df, min_baseline_age=args.min_baseline_age)

    with timer.track("write_variable_rules_sample_flow_table1"):
        if not globals().get("_PRIMARY8_OVERLAY_INSTALLED", False):
            fast_to_csv(variable_rules(), outdir / OUT_VARIABLE_RULES)
        sample_flow_table(long_df, ana, path_df, min_baseline_age=args.min_baseline_age).to_csv(outdir / OUT_SAMPLE_FLOW, index=False)
        export_characteristics_xlsx(ana, outdir / OUT_TABLE1)
        fast_to_csv(shock_count_distribution(ana), outdir / OUT_SHOCK_COUNT_DIST)

    imputed_main = impute_covariates(ana, m=args.imputations, seed=args.seed, mi_max_iter=args.mi_max_iter, mi_n_nearest_features=args.mi_n_nearest_features, timer=timer)
    imputed_path = impute_covariates(path_df, m=args.imputations, seed=args.seed + 3000, mi_max_iter=args.mi_max_iter, mi_n_nearest_features=args.mi_n_nearest_features, timer=timer)
    optimizer_seq = [x.strip() for x in args.optimizer_seq.split(",") if x.strip()]

    with timer.track("fit_main_summary_model123"):
        summary_results, summary_diag = fit_terms_model123(imputed_main, SUMMARY_TERMS, timer, optimizer_seq=optimizer_seq, maxfun=args.maxfun, singular_tol=args.singular_tol, label_prefix="summary_")
        fast_to_csv(summary_results, outdir / OUT_MAIN_SUMMARY)

    with timer.track("fit_main_component_model123"):
        component_results, component_diag = fit_terms_model123(imputed_main, COMPONENT_TERMS, timer, optimizer_seq=optimizer_seq, maxfun=args.maxfun, singular_tol=args.singular_tol, label_prefix="component_")
        fast_to_csv(component_results, outdir / OUT_MAIN_COMPONENT)

    with timer.track("write_meta_ready_model3"):
        if not globals().get("_PRIMARY8_OVERLAY_INSTALLED", False):
            fast_to_csv(make_meta_ready(summary_results, component_results), outdir / OUT_META_READY)

    with timer.track("fit_subgroup_summary_model3"):
        subgroup_results = fit_subgroup_model3(imputed_main, exposures=SUMMARY_TERMS, timer=timer, optimizer_seq=optimizer_seq, maxfun=args.maxfun, singular_tol=args.singular_tol)
        fast_to_csv(subgroup_results, outdir / OUT_SUBGROUP)

    with timer.track("fit_interaction_summary_model3"):
        interaction_results, interaction_diag = fit_interaction_summary_model3(imputed_main, exposures=SUMMARY_TERMS, timer=timer, optimizer_seq=optimizer_seq, maxfun=args.maxfun, singular_tol=args.singular_tol)
        fast_to_csv(interaction_results, outdir / OUT_INTERACTION)

    with timer.track("fit_shock_count_model123"):
        shock_count_results, shock_count_diag = fit_shock_count_model123(imputed_main, timer=timer, optimizer_seq=optimizer_seq, maxfun=args.maxfun, singular_tol=args.singular_tol)
        fast_to_csv(shock_count_results, outdir / OUT_SHOCK_COUNT_MODEL)
        if not globals().get("_PRIMARY8_OVERLAY_INSTALLED", False):
            fast_to_csv(make_shock_count_meta_ready(shock_count_results), outdir / OUT_SHOCK_COUNT_META)

    path_specs = [
        ("Path1_IndShock_to_SRH_to_Depression", "IndShock_main", "next_self_rated_health", "gaussian", "self_rated_health"),
        ("Path2_IndShock_to_ADL_to_Depression", "IndShock_main", "next_adl", "binomial", "adl"),
        ("Path3_SocShock_to_SocialActivityDecline_to_Depression", "SocShock_main", "next_social_activity_decline", "binomial", "social_activity_decline"),
    ]
    path_point = fit_path_models_mi(imputed_path, path_specs, timer=timer)
    if not globals().get("_PRIMARY8_OVERLAY_INSTALLED", False):
        fast_to_csv(path_point, outdir / OUT_PATH_SUMMARY)

    boot_df = bootstrap_paths_mi(imputed_path, path_specs, n_boot=args.bootstraps, seed=args.seed + 7000, timer=timer)
    path_boot = summarize_bootstrap(path_point, boot_df)
    fast_to_csv(path_boot, outdir / OUT_PATH_BOOT)

    notes = "\n".join([
        f"input={input_path}",
        f"output_dir={outdir}",
        "project=MHAS primary incident-shock expanded-social-shock pathway pipeline",
        "dataset=Gateway Harmonized MHAS Version D",
        "analysis_waves=wave2-wave5",
        f"baseline_age_restriction=baseline_age >= {args.min_baseline_age}",
        "baseline_age_source=Wave 1 r1agey when available; first nonmissing age used for refreshment/new-spouse respondents",
        "wave1_used_only_for_incident_shock_derivation=yes",
        "wave6_excluded_from_primary_analysis=yes; MHAS wave6 should be used only in optional sensitivity analyses if needed",
        "pathway_construction_waves=wave2-wave5; usable pathway exposure rows require next-wave outcomes and therefore are wave2-wave4",
        "shock_definition=previous wave status 0 and current wave status 1; child death is current deceased-child count > previous count; wave1 incident shocks set missing",
        "income_level_3=wave-specific tertiles based on HHwITOT; 1 low, 2 middle, 3 high",
        "income_decline_shock=previous income_level_3 > current income_level_3; included in SocShock_main",
        "living_alone_transition=previous living_alone=0 and current living_alone=1; included in SocShock_main",
        "reviewer_requested_analysis_1=extended subgroup/interaction by pre-shock household income, pre-shock social activity, and pre-shock living-alone status; continuous age retained within age_group subgroup/interaction models",
        "reviewer_requested_analysis_2=cumulative shock burden category models: 0, 1, 2, and >=3 shocks",
        "reviewer_requested_analysis_3=shock-count frequency distribution: 0, 1, 2, 3, and >=4 shocks",
        "education=RAEDUCL time-invariant copied to every person-wave",
        "main_model_backend=R/lme4::glmer",
        "main_model_model1=exposure + random intercept + optional wave fixed effects",
        "household_income_raw_not_in_main_analysis=yes; income_decline_shock included as social shock",
        "overall_shock_main=row_any_shock(IndShock_main, SocShock_main); observed-component rule: any observed 1 -> 1, otherwise observed rows -> 0, all components missing -> missing; SocShock_main includes income_decline_shock and living_alone_transition",
        "main_model_model2=model1 + age(per 10-year increase) + male + factor(education)",
        "main_model_model3=model2 + smoking + drinking; age coefficient retained as per-10-year effect after R-side scaling age/10",
        f"include_wave_fixed_effect={INCLUDE_WAVE_FIXED_EFFECT}",
        "self_rated_health_and_adl_removed_from_main_covariates=yes",
        "path_model_backend=GEE + bootstrap",
        "path_mediation_primary=counterfactual parametric g-computation on risk-difference scale",
        "path_mediation_secondary=coefficient-difference estimate on logit and y-standardized logit scales; no a*b product of log-ORs is used",
        "path_1=IndShock_main -> next self_rated_health -> next Depression_Risk",
        "path_2=IndShock_main -> next adl -> next Depression_Risk",
        "path_3=SocShock_main -> next social_activity_decline -> next Depression_Risk",
        "social_activity_variable=rWsocmn preferred; rWsocwk fallback; available from wave3 onward in MHAS",
        "smoking_variable_priority=rWsmoken/RwSMOKEN; fallback rWsmokev/RwSMOKEV only if current smoking is unavailable",
        "drinking_variable_priority=rWdrink/RwDRINK for MHAS",
        "traffic_accident_or_hospitalization_not_in_main_analysis=yes",
        "MI_for_covariates_in_main_and_path_models=age,male,education,smoking,drinking",
        "mediators_exposures_outcome_not_imputed=yes",
        f"imputations={args.imputations}",
        f"bootstraps={args.bootstraps}",
        f"optimizer_seq={args.optimizer_seq}",
        f"maxfun={args.maxfun}",
        f"singular_tol={args.singular_tol}",
        f"summary_diag_any_problem={diag_problem(summary_diag)}",
        f"component_diag_any_problem={diag_problem(component_diag)}",
        f"interaction_diag_any_problem={diag_problem(interaction_diag)}",
        f"shock_count_diag_any_problem={diag_problem(shock_count_diag)}",
        "",
        "Essential outputs only:",
        f"- {OUT_VARIABLE_RULES}",
        f"- {OUT_SAMPLE_FLOW}",
        f"- {OUT_TABLE1}",
        f"- {OUT_MAIN_SUMMARY}",
        f"- {OUT_MAIN_COMPONENT}",
        f"- {OUT_SUBGROUP}",
        f"- {OUT_INTERACTION}",
        f"- {OUT_SHOCK_COUNT_DIST}",
        f"- {OUT_SHOCK_COUNT_MODEL}",
        f"- {OUT_SHOCK_COUNT_META}",
        f"- {OUT_PATH_SUMMARY}",
        f"- {OUT_PATH_BOOT}",
        f"- {OUT_META_READY}",
        f"- {OUT_TIMING}",
    ]) + "\n"
    if not globals().get("_PRIMARY8_OVERLAY_INSTALLED", False):
        write_text(outdir / OUT_NOTES, notes)
        timer.to_frame().to_csv(outdir / OUT_TIMING, index=False)
    print("[Done] MHAS group-meeting incident-income-path pipeline completed.", flush=True)
    print(f"Results dir: {outdir}", flush=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="MHAS group-meeting incident shock + income decline + pathway pipeline")
    parser.add_argument("--input", type=str, default="H_MHAS_d.dta")
    parser.add_argument("--outdir", type=str, default="MHAS")
    parser.add_argument("--imputations", type=int, default=10)
    parser.add_argument("--min-baseline-age", type=float, default=MIN_BASELINE_AGE_DEFAULT)
    parser.add_argument("--bootstraps", type=int, default=200)
    parser.add_argument("--seed", type=int, default=20260423)
    parser.add_argument("--mi-max-iter", type=int, default=8)
    parser.add_argument("--mi-n-nearest-features", type=int, default=4)
    parser.add_argument("--optimizer-seq", type=str, default="bobyqa,Nelder_Mead,nloptwrap")
    parser.add_argument("--maxfun", type=int, default=200000)
    parser.add_argument("--singular-tol", type=float, default=1e-4)
    return parser.parse_args()


if __name__ == "__main__":
    run_pipeline(parse_args())
