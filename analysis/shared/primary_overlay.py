#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Install the revised eight-shock primary specification on the legacy cohort modules.

The uploaded cohort programs are retained as the modelling backend so the revised
primary analysis uses the legacy cohort-specific modelling backends while applying
the final eight-shock specification. Public-facing outputs are intentionally limited
to manuscript-relevant result files; audit, timing, and plotting artifacts are omitted.
"""
from __future__ import annotations

import gc
from pathlib import Path
from typing import List

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data_integrity import country_values
from .enhanced_mi import covariate_mi

from .project_config import (
    ALL_COMPONENTS_8,
    IND_COMPONENTS_8,
    LABOR_STATUS_SPECS,
    MARITAL_STATUS_SPECS,
    MODEL3_COVARIATES,
    SOC_COMPONENTS_8,
    SUMMARY_TERMS,
)


def _numeric(s) -> pd.Series:
    return pd.to_numeric(s, errors="coerce")


def _domain_from_components(d: pd.DataFrame, components: List[str]) -> pd.Series:
    """Conservative any-vs-none rule: any 1 -> 1; all resolved 0 -> 0; otherwise NA."""
    x = d[components].apply(pd.to_numeric, errors="coerce")
    any_one = x.eq(1).any(axis=1)
    all_resolved = x.notna().all(axis=1)
    out = pd.Series(np.nan, index=d.index, dtype="float32")
    out.loc[any_one] = 1.0
    out.loc[(~any_one) & all_resolved] = 0.0
    return out


def recompute_primary_domains(d: pd.DataFrame) -> pd.DataFrame:
    out = d.copy()
    out["IndShock_main"] = _domain_from_components(out, IND_COMPONENTS_8)
    out["SocShock_main"] = _domain_from_components(out, SOC_COMPONENTS_8)
    h = _numeric(out["IndShock_main"])
    s = _numeric(out["SocShock_main"])
    overall = pd.Series(np.nan, index=out.index, dtype="float32")
    overall.loc[h.eq(1) | s.eq(1)] = 1.0
    overall.loc[h.eq(0) & s.eq(0)] = 0.0
    out["OverallShock_main"] = overall
    return out


def domain_with_optional_spouse(d: pd.DataFrame) -> pd.Series:
    """Social domain for S3 where spouse cancer can be structurally not applicable.

    Structural N/A is ignored as an inapplicable component; true missing within an
    applicable spouse-risk set blocks a zero classification unless another event is 1.
    """
    regular = SOC_COMPONENTS_8
    x = d[regular].apply(pd.to_numeric, errors="coerce")
    spouse = _numeric(d.get("spouse_cancer", np.nan))
    applicability = _numeric(d["spouse_cancer_applicable"])
    applicable = applicability.eq(1)
    any_one = x.eq(1).any(axis=1) | spouse.eq(1)
    regular_all_zero = x.notna().all(axis=1) & x.eq(0).all(axis=1)
    spouse_resolved_zero_or_na = applicability.eq(0) | spouse.eq(0)
    out = pd.Series(np.nan, index=d.index, dtype="float32")
    out.loc[any_one] = 1.0
    out.loc[(~any_one) & regular_all_zero & spouse_resolved_zero_or_na] = 0.0
    return out


def _resolve_labor_col(mod, input_path: Path, cohort: str, wave: int) -> str:
    existing = mod.existing_columns(input_path)
    lower = {str(c).lower(): str(c) for c in existing}
    for pat in LABOR_STATUS_SPECS[cohort]["patterns"]:
        key = pat.format(w=wave).lower()
        if key in lower:
            return lower[key]
    raise KeyError(
        f"[{cohort}] labor-force status unavailable at wave {wave}; "
        f"tried {[p.format(w=wave) for p in LABOR_STATUS_SPECS[cohort]['patterns']]}"
    )


def _strict_incident(curr: pd.Series, prev: pd.Series, valid_pair: pd.Series) -> pd.Series:
    curr = _numeric(curr)
    prev = _numeric(prev)
    known = valid_pair & curr.isin([0, 1]) & prev.isin([0, 1])
    out = pd.Series(np.nan, index=curr.index, dtype="float32")
    out.loc[known] = 0.0
    out.loc[known & prev.eq(0) & curr.eq(1)] = 1.0
    return out


def apply_strict8_long(long_df: pd.DataFrame, cohort: str) -> pd.DataFrame:
    cohort = cohort.upper()
    labor = LABOR_STATUS_SPECS[cohort]
    marital = MARITAL_STATUS_SPECS[cohort]
    d = long_df.sort_values(["id", "wave"], kind="mergesort").copy()

    d["prev_inw_strict"] = d.groupby("id", sort=False)["inw"].shift(1)
    valid_pair = _numeric(d["inw"]).eq(1) & _numeric(d["prev_inw_strict"]).eq(1)
    d["valid_adjacent_interview_pair"] = valid_pair.astype("int8")

    # Health incident events; disease-specific component models later use prior disease-negative risk sets.
    for exposure, state in {
        "cancer": "cancer_state",
        "stroke": "stroke_state",
        "heart": "heart_state",
        "diabetes": "diabetes_state",
    }.items():
        prev_col = f"prev_{state}"
        if prev_col not in d:
            d[prev_col] = d.groupby("id", sort=False)[state].shift(1)
        d[exposure] = _strict_incident(d[state], d[prev_col], valid_pair)
        prev = _numeric(d[prev_col]); curr = _numeric(d[state])
        d[f"risk_{exposure}"] = (valid_pair & prev.eq(0) & curr.isin([0, 1])).astype("int8")

    # Strict marital disruption: previously married/stable partnered -> separated/divorced/widowed.
    if "marital_status_strict_raw" not in d:
        d["marital_status_strict_raw"] = _numeric(d.get("mstat_raw", np.nan))
    d["prev_marital_status_strict_raw"] = d.groupby("id", sort=False)["marital_status_strict_raw"].shift(1)
    prev_m = _numeric(d["prev_marital_status_strict_raw"])
    curr_m = _numeric(d["marital_status_strict_raw"])
    m_known = valid_pair & prev_m.isin(marital["known_codes"]) & curr_m.isin(marital["known_codes"])
    prev_partnered = prev_m.isin(marital["partnered_codes"])
    curr_partnered = curr_m.isin(marital["partnered_codes"])
    marital_event = pd.Series(np.nan, index=d.index, dtype="float32")
    marital_event.loc[m_known] = 0.0
    marital_event.loc[m_known & prev_partnered & curr_m.isin(marital["adverse_codes"])] = 1.0
    d["marital_shock"] = marital_event
    d["risk_marital_shock"] = (m_known & prev_partnered).astype("int8")
    d["partnered_prev"] = np.where(prev_m.isin(marital["known_codes"]), prev_partnered.astype(float), np.nan)
    d["partnered_current"] = np.where(curr_m.isin(marital["known_codes"]), curr_partnered.astype(float), np.nan)

    # Work shock: previous active work -> current explicit unemployment only.
    if "labor_status_strict" not in d:
        raise KeyError(f"[{cohort}] labor_status_strict missing after wave-frame construction")
    d["prev_labor_status_strict"] = d.groupby("id", sort=False)["labor_status_strict"].shift(1)
    prev_l = _numeric(d["prev_labor_status_strict"])
    curr_l = _numeric(d["labor_status_strict"])
    l_known = valid_pair & prev_l.isin(labor["known_codes"]) & curr_l.isin(labor["known_codes"])
    prev_working = prev_l.isin(labor["working_codes"])
    curr_unemployed = curr_l.isin(labor["unemployed_codes"])
    job = pd.Series(np.nan, index=d.index, dtype="float32")
    job.loc[l_known] = 0.0
    job.loc[l_known & prev_working & curr_unemployed] = 1.0
    d["job_exit"] = job
    d["risk_job_exit"] = (l_known & prev_working).astype("int8")

    # Child death is cohort-specific in the legacy builder; only enforce valid adjacent interviews.
    child = _numeric(d.get("child_death", np.nan))
    d.loc[~valid_pair, "child_death"] = np.nan
    d["risk_child_death"] = (valid_pair & child.isin([0, 1])).astype("int8")

    # Living-alone transition; only previous non-alone respondents are in the component risk set.
    if "prev_living_alone" not in d:
        d["prev_living_alone"] = d.groupby("id", sort=False)["living_alone"].shift(1)
    prev_live = _numeric(d["prev_living_alone"]); curr_live = _numeric(d["living_alone"])
    live_known = valid_pair & prev_live.isin([0, 1]) & curr_live.isin([0, 1])
    live = pd.Series(np.nan, index=d.index, dtype="float32")
    live.loc[live_known] = 0.0
    live.loc[live_known & prev_live.eq(0) & curr_live.eq(1)] = 1.0
    d["living_alone_transition"] = live
    d["risk_living_alone_transition"] = (live_known & prev_live.eq(0)).astype("int8")

    # Keep the two non-primary shocks available for S2/S3.
    if "income_decline_shock" in d:
        d.loc[~valid_pair, "income_decline_shock"] = np.nan
        d["risk_income_decline_shock"] = valid_pair.astype("int8")

    # Strict spouse-cancer variable for S3. No partner is structural N/A, never recoded as a no-event.
    if "spouse_cancer_state" in d:
        if "prev_spouse_cancer_state" not in d:
            d["prev_spouse_cancer_state"] = d.groupby("id", sort=False)["spouse_cancer_state"].shift(1)
        prev_sc = _numeric(d["prev_spouse_cancer_state"])
        curr_sc = _numeric(d["spouse_cancer_state"])
        applicable = valid_pair & prev_partnered & curr_partnered
        sc = pd.Series(np.nan, index=d.index, dtype="float32")
        observed = applicable & prev_sc.isin([0, 1]) & curr_sc.isin([0, 1])
        sc.loc[observed] = 0.0
        sc.loc[observed & prev_sc.eq(0) & curr_sc.eq(1)] = 1.0
        d["spouse_cancer"] = sc
        applicability = pd.Series(np.nan, index=d.index, dtype="float32")
        known_nonpartner = valid_pair & ((prev_m.isin(marital["known_codes"]) & ~prev_partnered) | (curr_m.isin(marital["known_codes"]) & ~curr_partnered))
        applicability.loc[known_nonpartner] = 0.
        applicability.loc[applicable] = 1.
        d["spouse_cancer_applicable"] = applicability
        d["risk_spouse_cancer"] = applicable.astype("int8")

    # New eight-shock primary domains.
    d = recompute_primary_domains(d)
    return d


def add_strict_shock_count_variables(df: pd.DataFrame) -> pd.DataFrame:
    """Primary cumulative burden: any missing retained shock -> burden missing/excluded."""
    out = df.copy()
    comp = out[ALL_COMPONENTS_8].apply(pd.to_numeric, errors="coerce")
    complete = comp.notna().all(axis=1)
    out["shock_count_missing_components"] = comp.isna().sum(axis=1).astype("float32")
    out["shock_count_all"] = np.nan
    out["shock_count_cat"] = np.nan
    out["shock_count_trend"] = np.nan
    out["shock_count_plot"] = np.nan
    if complete.any():
        counts = comp.loc[complete].sum(axis=1).astype(float)
        out.loc[complete, "shock_count_all"] = counts
        out.loc[complete, "shock_count_cat"] = counts.clip(upper=3)
        out.loc[complete, "shock_count_trend"] = counts.clip(upper=3)
        out.loc[complete, "shock_count_plot"] = counts.clip(upper=4)
    return out


def variable_rules_8() -> pd.DataFrame:
    rows = [
        ("Primary shock set", "8 retained components", "Cancer, stroke, heart disease, diabetes, marital disruption, child death, work-to-unemployment, transition to living alone"),
        ("Incident timing", "consecutive interviews", "Shock is first recorded between the immediately preceding and current interview; both interviews must be observed"),
        ("cancer", "harmonized disease state", "1 only for previous 0 -> current 1; component model risk set requires previous disease=0"),
        ("stroke", "harmonized disease state", "1 only for previous 0 -> current 1; component model risk set requires previous disease=0"),
        ("heart", "harmonized disease state", "1 only for previous 0 -> current 1; component model risk set requires previous disease=0"),
        ("diabetes", "harmonized disease state", "1 only for previous 0 -> current 1; component model risk set requires previous disease=0"),
        ("marital_shock", "harmonized marital status", "Previous married/stable partnered -> current separated/divorced/widowed; component model restricted to previous partnered risk set"),
        ("child_death", "cohort-specific harmonized child-death measure", "Cohort-specific incident construction retained; requires valid adjacent interviews"),
        ("job_exit", "harmonized labor-force status", "Previous active work -> current explicit unemployment only; retirement/partial retirement/disability/homemaking/not-in-labor-force/never-worked are not shocks"),
        ("living_alone_transition", "household size", "Previous household size >=2 -> current household size=1; component model restricted to previous non-alone risk set"),
        ("IndShock_main", "4 health components", "Any known component=1 -> 1; 0 only when all four are resolved as 0; otherwise missing"),
        ("SocShock_main", "4 social components", "Any known component=1 -> 1; 0 only when all four are resolved as 0; otherwise missing"),
        ("OverallShock_main", "health + social domains", "1 if either domain=1; 0 only if both domains=0; otherwise missing"),
        ("Cumulative burden", "all 8 retained components", "Calculated only when all eight shock components are determinable; any missing component makes the person-wave burden missing and excludes that person-wave from burden analysis"),
        ("Primary covariate missingness", "age/sex/education/smoking/drinking", "Multiple imputation; exposure and depressive outcome are not imputed in the primary analysis"),
    ]
    return pd.DataFrame(rows, columns=["analysis_variable", "raw_variable", "definition"])


def sample_flow_8(mod, long_df: pd.DataFrame, analysis_df: pd.DataFrame, path_df: pd.DataFrame, min_baseline_age: float) -> pd.DataFrame:
    rows = []
    rows.append({"step": "wide_persons_input", "n": int(long_df["id"].nunique())})
    rows.append({"step": "long_rows_all_loaded_waves", "n": int(len(long_df))})
    age_ok = _numeric(long_df["baseline_age"]).ge(min_baseline_age)
    rows.append({"step": f"persons_baseline_age_ge_{int(min_baseline_age)}", "n": int(long_df.loc[age_ok, "id"].nunique())})
    tmp = long_df[long_df["wave"].isin(mod.ANALYSIS_WAVES) & age_ok].copy()
    rows.append({"step": "rows_prespecified_analysis_waves_age_eligible", "n": int(len(tmp))})
    tmp = tmp[_numeric(tmp["inw"]).eq(1)]
    rows.append({"step": "rows_current_wave_responded", "n": int(len(tmp))})
    tmp = tmp[tmp[mod.OUTCOME].notna()]
    rows.append({"step": "rows_nonmissing_outcome", "n": int(len(tmp))})
    tmp = tmp[_numeric(tmp["valid_adjacent_interview_pair"]).eq(1)]
    rows.append({"step": "rows_valid_immediate_previous_interview", "n": int(len(tmp))})
    all8 = tmp.dropna(subset=ALL_COMPONENTS_8)
    rows.append({"step": "rows_all_8_shocks_determinable_for_cumulative_burden", "n": int(len(all8))})
    rows.append({"step": "rows_in_primary_model_base_dataset", "n": int(len(analysis_df))})
    rows.append({"step": "persons_in_primary_model_base_dataset", "n": int(analysis_df["id"].nunique())})
    rows.append({"step": "rows_in_path_dataset", "n": int(len(path_df))})
    rows.append({"step": "persons_in_path_dataset", "n": int(path_df["id"].nunique()) if len(path_df) else 0})
    return pd.DataFrame(rows)


def export_characteristics_xlsx_8(mod, df: pd.DataFrame, out_path: Path) -> None:
    specs = [
        ("Outcome", None, None),
        ("Depression_Risk, n (%)", mod.OUTCOME, "binary"),
        ("Health shock components", None, None),
        ("Cancer, n (%)", "cancer", "binary"),
        ("Stroke, n (%)", "stroke", "binary"),
        ("Heart disease, n (%)", "heart", "binary"),
        ("Diabetes, n (%)", "diabetes", "binary"),
        ("Social shock components", None, None),
        ("Marital disruption, n (%)", "marital_shock", "binary"),
        ("Death of a child, n (%)", "child_death", "binary"),
        ("Work-to-unemployment, n (%)", "job_exit", "binary"),
        ("Transition to living alone, n (%)", "living_alone_transition", "binary"),
        ("Domains", None, None),
        ("Overall shock, n (%)", "OverallShock_main", "binary"),
        ("Health-related shock, n (%)", "IndShock_main", "binary"),
        ("Social-disruption shock, n (%)", "SocShock_main", "binary"),
        ("Covariates", None, None),
        ("Baseline age, mean (SD)", "baseline_age", "cont"),
        ("Age, mean (SD)", "age", "cont"),
        ("Age >60, n (%)", "age_group", "binary"),
        ("Male, n (%)", "male", "binary"),
        ("Education, mean (SD)", "education", "cont"),
        ("Smoking, n (%)", "smoking", "binary"),
        ("Drinking, n (%)", "drinking", "binary"),
    ]
    waves = list(mod.ANALYSIS_WAVES)
    wb = Workbook(); ws = wb.active; ws.title = "Characteristics"
    wave_ns = {int(w): int((_numeric(df["wave"]) == int(w)).sum()) for w in waves}
    header = ["Characteristics"] + [f"Wave{w}\n(N={wave_ns[int(w)]})" for w in waves] + [f"Overall\n(N={len(df)})", "P-Value"]
    ws.append(header)
    section_rows = []
    for label, var, typ in specs:
        if var is None:
            ws.append([label] + [""] * (len(header)-1)); section_rows.append(ws.max_row); continue
        row = [label]
        for w in waves:
            s = _numeric(df.loc[_numeric(df["wave"]).eq(int(w)), var])
            if typ == "binary":
                n = int(s.notna().sum()); e = int(s.eq(1).sum()); row.append("NA" if n == 0 else f"{e} ({100*e/n:.1f})")
            else:
                row.append("NA" if s.notna().sum() == 0 else f"{s.mean():.1f} ({s.std(ddof=1):.1f})")
        s = _numeric(df[var])
        if typ == "binary":
            n = int(s.notna().sum()); e = int(s.eq(1).sum()); row.append("NA" if n == 0 else f"{e} ({100*e/n:.1f})")
            try: p = mod._binary_pvalue(df, var)
            except Exception: p = np.nan
        else:
            row.append("NA" if s.notna().sum() == 0 else f"{s.mean():.1f} ({s.std(ddof=1):.1f})")
            try: p = mod._cont_pvalue(df, var)
            except Exception: p = np.nan
        row.append("" if pd.isna(p) else ("<0.001" if p < .001 else f"{p:.3f}")); ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78"); section_fill = PatternFill("solid", fgColor="D9EAF7"); thin = Side(style="thin", color="BFBFBF")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = header_fill; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)
    for r in range(2, ws.max_row+1):
        for c in range(1, len(header)+1):
            cell = ws.cell(r,c); cell.alignment = Alignment(horizontal="left" if c==1 else "center", vertical="center", wrap_text=True); cell.border = Border(bottom=thin)
            if r in section_rows: cell.fill = section_fill; cell.font = Font(bold=True)
    for c in range(1, len(header)+1): ws.column_dimensions[get_column_letter(c)].width = 36 if c==1 else 16
    ws.freeze_panes = "B2"; wb.save(out_path)


def install_primary8_overlay(mod, cohort: str):
    """Monkey-patch one uploaded legacy cohort module with the final 8-shock specification."""
    cohort = cohort.upper()
    if getattr(mod, "_PRIMARY8_OVERLAY_INSTALLED", False):
        return mod

    # Store originals before patching.
    orig_cast_small_dtypes = mod.cast_small_dtypes
    orig_build_raw_column_map = mod.build_raw_column_map
    orig_make_wave_frame = mod.make_wave_frame
    orig_build_long_panel = mod.build_long_panel
    orig_prepare_main_dataset = mod.prepare_main_dataset
    orig_prepare_path_dataset = mod.prepare_path_dataset
    orig_fit_one = mod.fit_one_exposure_all_imputations
    orig_run_pipeline = mod.run_pipeline

    mod.IND_COMPONENTS = list(IND_COMPONENTS_8)
    mod.SOC_COMPONENTS = list(SOC_COMPONENTS_8)
    mod.COMPONENT_TERMS = list(ALL_COMPONENTS_8)
    mod.SUMMARY_TERMS = list(SUMMARY_TERMS)

    def patched_build_raw_column_map(path: Path):
        m = orig_build_raw_column_map(path)
        mod._RAW_ID_COLUMN = m["id"]
        existing = mod.existing_columns(path)
        lower = {str(c).lower(): str(c) for c in existing}
        for w in mod.ALL_WAVES:
            found = None
            for pat in LABOR_STATUS_SPECS[cohort]["patterns"]:
                key = pat.format(w=w).lower()
                if key in lower:
                    found = lower[key]; break
            if found is None:
                raise KeyError(f"[{cohort}] cannot resolve labor-force status for wave {w}")
            m["waves"][w]["labor_status_strict"] = found
            if cohort == "SHARE":
                candidates = ["country", "racountry", "ra_country", "countryid", f"r{w}country"]
                country_col = next((lower[c] for c in candidates if c in lower), None)
                if country_col is None:
                    raise KeyError(f"[SHARE] country variable required; tried {candidates}")
                m["waves"][w]["country"] = country_col
        return m

    def patched_make_wave_frame(df: pd.DataFrame, w: int, rawmap: dict):
        out = orig_make_wave_frame(df, w, rawmap)
        wm = rawmap["waves"][w]
        out["labor_status_strict"] = _numeric(df[wm["labor_status_strict"]]).to_numpy()
        out["marital_status_strict_raw"] = _numeric(df[wm["mstat"]]).to_numpy()
        if cohort == "SHARE":
            out["country"] = country_values(df[wm["country"]]).to_numpy()
            if out.loc[_numeric(out["inw"]).eq(1), "country"].isna().any():
                raise ValueError("[SHARE] missing/invalid country among respondents; resolve source mapping before fitting")
        # Preserve raw spouse state. Applicability is handled explicitly below;
        # legacy zero-filling outside its narrower partnered codes is not reused.
        sc_col = wm.get("spouse_cancer_state")
        if sc_col is not None:
            raw_sc = _numeric(df[sc_col])
            out["spouse_cancer_state"] = raw_sc.where(raw_sc.isin([0, 1])).to_numpy()
        return out

    def patched_cast_small_dtypes(frame):
        idcol = getattr(mod, "_RAW_ID_COLUMN", None)
        if idcol is None or idcol not in frame:
            return orig_cast_small_dtypes(frame)
        saved_id = frame[idcol].copy()
        result = orig_cast_small_dtypes(frame.drop(columns=[idcol]).copy())
        result[idcol] = saved_id
        return result[frame.columns]

    def patched_impute_covariates(base_df, m, seed, mi_max_iter, mi_n_nearest_features, timer):
        # Accept the legacy nearest-features option for CLI compatibility; the
        # enhanced model deliberately uses all requested auxiliary predictors.
        with timer.track("enhanced_multiple_imputation_covariates"):
            imps, _ = covariate_mi(base_df, m, seed, mi_max_iter)
        return imps

    def patched_build_long_panel(df_wide: pd.DataFrame, rawmap: dict):
        d = orig_build_long_panel(df_wide, rawmap)
        return apply_strict8_long(d, cohort)

    def patched_prepare_main_dataset(long_df: pd.DataFrame, min_baseline_age: float = 50.0):
        d = orig_prepare_main_dataset(long_df, min_baseline_age=min_baseline_age)
        d = d[_numeric(d["valid_adjacent_interview_pair"]).eq(1)].copy()
        return d.sort_values(["id", "wave"], kind="mergesort").reset_index(drop=True)

    def patched_prepare_path_dataset(long_df: pd.DataFrame, min_baseline_age: float = 50.0):
        d = orig_prepare_path_dataset(long_df, min_baseline_age=min_baseline_age)
        if "valid_adjacent_interview_pair" in d:
            d = d[_numeric(d["valid_adjacent_interview_pair"]).eq(1)].copy()
        return d.reset_index(drop=True)

    def patched_fit_one(imputed_list, exposure, model_id, timer, optimizer_seq, maxfun, singular_tol, omit=None, label_prefix=""):
        risk_col = f"risk_{exposure}"
        if exposure in ALL_COMPONENTS_8 and any(risk_col in x.columns for x in imputed_list):
            filt = []
            for x in imputed_list:
                if risk_col in x:
                    filt.append(x[_numeric(x[risk_col]).eq(1)].copy())
                else:
                    filt.append(x)
            return orig_fit_one(filt, exposure, model_id, timer, optimizer_seq, maxfun, singular_tol, omit=omit, label_prefix=label_prefix)
        return orig_fit_one(imputed_list, exposure, model_id, timer, optimizer_seq, maxfun, singular_tol, omit=omit, label_prefix=label_prefix)

    def patched_sample_flow(long_df, analysis_df, path_df, min_baseline_age=50.0):
        return sample_flow_8(mod, long_df, analysis_df, path_df, min_baseline_age)

    def patched_export_characteristics(df, out_path):
        return export_characteristics_xlsx_8(mod, df, out_path)

    def patched_run_pipeline(args):
        outdir = Path(args.outdir).resolve()
        outdir.mkdir(parents=True, exist_ok=True)
        mod._REVISION_OUTDIR = outdir
        orig_run_pipeline(args)

        # The public repository keeps only manuscript-relevant analysis outputs.
        # Remove duplicated/meta-ready, diagnostic, note, and timing files emitted
        # by the legacy backend. No audit files are written by this overlay.
        discard = [
            getattr(mod, "OUT_VARIABLE_RULES", None),
            getattr(mod, "OUT_PATH_SUMMARY", None),
            getattr(mod, "OUT_META_READY", None),
            getattr(mod, "OUT_SHOCK_COUNT_META", None),
            getattr(mod, "OUT_NOTES", None),
            getattr(mod, "OUT_TIMING", None),
        ]
        for name in discard:
            if name:
                path = outdir / name
                if path.exists():
                    path.unlink()

    mod._COUNTRY_FE_ENABLED = cohort == "SHARE"
    mod.cast_small_dtypes = patched_cast_small_dtypes
    mod.impute_covariates = patched_impute_covariates
    mod.build_raw_column_map = patched_build_raw_column_map
    mod.make_wave_frame = patched_make_wave_frame
    mod.build_long_panel = patched_build_long_panel
    mod.prepare_main_dataset = patched_prepare_main_dataset
    mod.prepare_path_dataset = patched_prepare_path_dataset
    mod.add_shock_count_variables = add_strict_shock_count_variables
    mod.variable_rules = variable_rules_8
    mod.sample_flow_table = patched_sample_flow
    mod.export_characteristics_xlsx = patched_export_characteristics
    mod.fit_one_exposure_all_imputations = patched_fit_one
    mod.run_pipeline = patched_run_pipeline
    mod.recompute_primary_domains = recompute_primary_domains
    mod.domain_with_optional_spouse = domain_with_optional_spouse
    mod.ALL_COMPONENTS_8 = list(ALL_COMPONENTS_8)
    mod._PRIMARY8_OVERLAY_INSTALLED = True
    gc.collect()
    return mod
